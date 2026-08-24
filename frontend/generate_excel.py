"""Genere le classeur d'import de l'historique VaR marche.

Deux feuilles, au format long attendu par `POST /market/upload-var-history`
(backend/app/market/routes.py) : une ligne = un titre a une date. Les
caracteristiques statiques alimentent la position (premiere occurrence de
l'identifiant), les couples (Date, Prix) alimentent l'historique.

Les feuilles « Historique Prix Obligations » et « Historique Cours Actions »
ont ete supprimees : elles portaient l'ancien format a quatre feuilles, que
l'import ne lit plus depuis que l'historique tient dans la feuille du titre.
"""

import openpyxl

wb = openpyxl.Workbook()

# Remove default sheet
default_sheet = wb.active
wb.remove(default_sheet)

# Obligations — positions + historique dans le meme onglet
ws_bonds = wb.create_sheet("Obligations")
bond_headers = [
    "ID Titre",
    "Emetteur",
    "Devise",
    "Valeur nominale unitaire",
    "Coupon (%)",
    "Fréquence de paiement des intérêts",
    "Date d'émission",
    "Date d'échéance",
    "quantités",
    "Date",
    "Prix de marché (%)",
]
ws_bonds.append(bond_headers)

# Actions — positions + historique dans le meme onglet
ws_equities = wb.create_sheet("Actions")
equity_headers = [
    "ID Instrument",
    "Émetteur / Société",
    "Secteur",
    "Quantité",
    "Date",
    "Cours de clôture",
]
ws_equities.append(equity_headers)

wb.save("Modele_Import_Risque_Marche.xlsx")
