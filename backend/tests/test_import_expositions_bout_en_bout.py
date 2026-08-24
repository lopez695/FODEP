"""Le modele d'import livre doit se reimporter tel quel.

C'est le seul controle qui vaut sur un import : produire le classeur que
l'application propose au telechargement, puis le lui rendre. Les tests
existants verifiaient que le modele portait les bonnes colonnes et que le
validateur reconnaissait ses feuilles -- aucun ne parcourait le chemin
jusqu'au bout.

Il manquait donc ceci : `_parse_workbook` passait a `collecter_identifications`
la liste rendue par `_read_sheet_rows`, qui est faite de couples (numero de
ligne Excel, valeurs). La collecte, elle, attend des dictionnaires. Tout import
d'expositions echouait sur `'tuple' object has no attribute 'get'` -- une
erreur 500, sur chaque fichier, quel qu'il soit.
"""

from __future__ import annotations

import warnings

import pytest

from database.services.excel_import_service import ExcelImportService


@pytest.fixture(scope="module")
def modele() -> bytes:
    service = ExcelImportService.__new__(ExcelImportService)
    with warnings.catch_warnings():
        # openpyxl previent qu'il ne sait pas relire l'extension x14 des
        # validations : sans effet sur les valeurs lues ici.
        warnings.simplefilter("ignore")
        return ExcelImportService.build_template_workbook(service)


def test_le_modele_livre_traverse_l_import_sans_echouer(modele):
    """La lecture du classeur va jusqu'aux enregistrements d'exposition.

    Le controle s'arrete a l'analyse : ecrire en base ferait entrer dans le
    portefeuille les dix-huit lignes d'exemple du modele.
    """

    from io import BytesIO

    from openpyxl import load_workbook

    from database.services.excel_import_service import ImportProfiler

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        classeur = load_workbook(BytesIO(modele), data_only=True, read_only=True)
    try:
        service = ExcelImportService.__new__(ExcelImportService)
        analyse = service._parse_workbook(classeur, profiler=ImportProfiler(source_name="modele livre"))
    finally:
        classeur.close()

    assert analyse.rows_read, "aucune ligne lue dans le modele livre"
    assert analyse.exposure_records, "aucune exposition reconstituee"
    assert not analyse.errors, f"lignes rejetees : {analyse.errors[:3]}"


def test_l_import_reprend_l_identification_des_contreparties(modele):
    """Les colonnes qui decrivent la contrepartie voyagent avec le classeur.

    C'est l'appel qui plantait. Le verifier ici, et pas seulement sur des
    dictionnaires fabriques a la main comme le fait
    `test_identification_import`, garde le raccord entre les deux.
    """

    from io import BytesIO

    from openpyxl import load_workbook

    from database.services.excel_import_service import ImportProfiler

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        classeur = load_workbook(BytesIO(modele), data_only=True, read_only=True)
    try:
        service = ExcelImportService.__new__(ExcelImportService)
        analyse = service._parse_workbook(classeur, profiler=ImportProfiler(source_name="modele livre"))
    finally:
        classeur.close()

    # Le type suffit : le modele livre peut ne renseigner aucune de ces
    # colonnes sur ses lignes d'exemple, mais la collecte doit avoir tourne.
    assert isinstance(analyse.identifications, dict)


def test_les_deux_declarations_de_colonnes_ne_divergent_pas():
    """La spec du validateur et celle du modele decrivent le meme classeur.

    `excel_repository.py` le dit deja en commentaire : sa liste « doit rester
    identique, colonne par colonne, a la spec du validateur ». Rien ne le
    verifiait, et les deux ont diverge -- six colonnes d'identification de la
    contrepartie declarees au validateur, lues a l'import, mais absentes des
    listes qui construisent le modele livre. Aucun classeur ne pouvait donc les
    porter, et l'EP29, l'EP30, l'EP32, l'EP38 et l'EP39 partaient sans elles.
    """

    from app.core.excel_repository import (
        EXPECTED_COLUMNS_BY_SHEET,
        OPTIONAL_COLUMNS_BY_SHEET,
    )
    from app.validators.excel_import_validator import IMPORT_SHEET_SPECS

    ecarts: list[str] = []
    for spec in IMPORT_SHEET_SPECS:
        cote_modele = set(EXPECTED_COLUMNS_BY_SHEET.get(spec.name, ())) | set(
            OPTIONAL_COLUMNS_BY_SHEET.get(spec.name, ())
        )
        cote_validateur = set(spec.import_columns)
        for colonne in sorted(cote_validateur - cote_modele):
            ecarts.append(f"{spec.name} : « {colonne} » lue a l'import, absente du modele")
        for colonne in sorted(cote_modele - cote_validateur):
            ecarts.append(f"{spec.name} : « {colonne} » portee par le modele, ignoree a l'import")

    assert not ecarts, "\n".join(ecarts)


def test_l_identification_de_la_contrepartie_traverse_le_classeur():
    """Une valeur portee dans le classeur doit ressortir de l'analyse.

    Le controle va du fichier jusqu'aux identifications collectees : c'est le
    chemin qu'empruntent le numero Centrale des risques et le groupe de clients
    lies, dont dependent cinq etats du FODEP.
    """

    from io import BytesIO

    from openpyxl import load_workbook

    from database.services.excel_import_service import ImportProfiler

    service = ExcelImportService.__new__(ExcelImportService)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        classeur = load_workbook(BytesIO(ExcelImportService.build_template_workbook(service)))
    feuille = classeur["Template données"]
    entetes = {
        str(feuille.cell(row=1, column=colonne).value).strip(): colonne
        for colonne in range(1, feuille.max_column + 1)
        if feuille.cell(row=1, column=colonne).value is not None
    }
    for colonne, valeur in (
        ("N_Centrale_risques", "CR-9001"),
        ("Groupe_clients_lies", "Groupe d'essai"),
        ("Secteur_activite", "Industrie"),
    ):
        feuille.cell(row=2, column=entetes[colonne], value=valeur)
    nom = feuille.cell(row=2, column=entetes["Contrepartie"]).value

    tampon = BytesIO()
    classeur.save(tampon)
    classeur.close()

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        relu = load_workbook(BytesIO(tampon.getvalue()), data_only=True, read_only=True)
    try:
        analyse = service._parse_workbook(
            relu, profiler=ImportProfiler(source_name="essai identification")
        )
    finally:
        relu.close()

    assert nom in analyse.identifications, (
        f"la contrepartie « {nom} » n'a pas ete identifiee : "
        f"{list(analyse.identifications)[:3]}"
    )
    porte = analyse.identifications[nom]
    assert porte.get("N_Centrale_risques") == "CR-9001"
    assert porte.get("Groupe_clients_lies") == "Groupe d'essai"
