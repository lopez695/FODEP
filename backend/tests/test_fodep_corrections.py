"""Ce que l'export corrige : homonymes, EP39, EP31, dérivés dans le levier.

Ces essais ne touchent ni la base ni le portefeuille : ils appellent les
fonctions de remplissage sur le modèle vierge, avec des données posées ici.
C'est ce qui permet de les lire comme des énoncés — « deux homonymes font deux
risques », « l'EP39 ne déclare que ce qui dépasse 5 % » — plutôt que comme des
comparaisons de totaux dépendant du portefeuille du jour.
"""

from __future__ import annotations

from types import SimpleNamespace
from unittest import mock

import pytest
from openpyxl import load_workbook

from app.rapports.fodep import agregation
from app.rapports.fodep.disposition import CHEMIN_MODELE, indexer_codes_dispru
from app.rapports.fodep.reserves import NatureReserve
from app.rapports.fodep.service import (
    COLONNE_C,
    COLONNE_TOTAL_EP39,
    SEUIL_EP39,
    _agreger_par_contrepartie,
    _bloc_derives,
    _ecarts_de_produit_brut,
    _exposition_totale,
    _remplir_ep33,
    _remplir_ep39,
)


@pytest.fixture
def classeur():
    fichier = load_workbook(CHEMIN_MODELE)
    yield fichier
    fichier.close()


def _exposition(identifiant: str, nom: str, bilan: float) -> dict:
    """Une ligne d'exposition telle que le dépôt la rend."""

    return {
        "counterparty_id": identifiant,
        "counterparty_name": nom,
        "country": "Côte d'Ivoire",
        "category_raw": "Entreprises",
        "currency": "XOF",
        "ead_bilan_amount": bilan,
        "ead_hb_ccf_amount": 0.0,
        "gross_amount": bilan,
        "rwa": bilan,
        "residual_maturity_months": 12,
        "prudential_type": "s",
    }


# ─── Agregation par identifiant ───────────────────────────────────────────


def test_deux_homonymes_font_deux_risques():
    """Le portefeuille porte 162 noms partages : les additionner gonflait tout.

    La division des risques declarait un risque unique la ou il y a deux
    clients, et l'EP38 comptait deux fois les concours d'une partie liee
    homonyme d'une autre.
    """

    groupes = _agreger_par_contrepartie([
        _exposition("EXP-0001", "M. KOUAME Adjoua", 1_000_000_000),
        _exposition("EXP-0002", "M. KOUAME Adjoua", 4_000_000_000),
    ])

    assert set(groupes) == {"EXP-0001", "EXP-0002"}
    assert _exposition_totale(groupes["EXP-0001"]) == 1_000_000_000
    assert _exposition_totale(groupes["EXP-0002"]) == 4_000_000_000
    # Le nom reste dans l'agregat : c'est lui que le formulaire imprime.
    assert groupes["EXP-0001"]["nom"] == "M. KOUAME Adjoua"


def test_une_exposition_sans_identifiant_retombe_sur_son_nom():
    """Une base ancienne ne doit pas perdre ses expositions."""

    exposition = _exposition("", "Société Sans Identifiant", 500_000_000)
    groupes = _agreger_par_contrepartie([exposition])
    assert list(groupes) == ["Société Sans Identifiant"]


# ─── EP39 : les montants, et le seuil imprime dans son titre ──────────────


def _groupes_parties_liees() -> tuple[dict, list[dict]]:
    groupes = _agreger_par_contrepartie([
        _exposition("EXP-0001", "M. DIALLO Fatou", 9_000_000_000),
        _exposition("EXP-0002", "Mme TRAORÉ Ali", 100_000_000),
    ])
    parties = [
        {"identifiant": "EXP-0001", "nom": "M. DIALLO Fatou", "categorie": "actionnaire"},
        {"identifiant": "EXP-0002", "nom": "Mme TRAORÉ Ali", "categorie": "cadre"},
    ]
    return groupes, parties


def test_ep39_ne_declare_que_ce_qui_depasse_cinq_pour_cent(classeur):
    """Le titre de l'etat fixe le seuil : 5 % des fonds propres effectifs.

    L'export declarait toute partie liee, et une croix au lieu d'un montant :
    la colonne TOTAL partait donc a zero en contradiction avec l'EP38.
    """

    groupes, parties = _groupes_parties_liees()
    fonds_propres = 100_000_000_000.0
    _remplir_ep39(classeur, groupes, parties, fonds_propres)

    feuille = classeur["EP39"]
    lignes = indexer_codes_dispru(feuille)
    codes = sorted(code for code in lignes if code.startswith("PR"))
    premiere, total = lignes[codes[0]], lignes[codes[-1]]

    # 9 Md dépassent 5 % de 100 Md ; 100 M ne les atteignent pas.
    assert feuille.cell(row=premiere, column=2).value == "M. DIALLO Fatou"
    assert feuille.cell(row=premiere, column=COLONNE_TOTAL_EP39).value == 9_000
    assert feuille.cell(row=total, column=COLONNE_TOTAL_EP39).value == 9_000
    noms = {
        feuille.cell(row=lignes[code], column=2).value
        for code in codes[:-1]
    }
    assert "Mme TRAORÉ Ali" not in noms


def test_ep39_porte_le_montant_dans_la_colonne_de_la_categorie(classeur):
    groupes, parties = _groupes_parties_liees()
    _remplir_ep39(classeur, groupes, parties, 100_000_000_000.0)

    feuille = classeur["EP39"]
    lignes = indexer_codes_dispru(feuille)
    premiere = lignes[sorted(code for code in lignes if code.startswith("PR"))[0]]
    # « Actionnaires détenant individuellement au moins 10 % » : colonne C.
    assert feuille.cell(row=premiere, column=COLONNE_C).value == 9_000


def test_ep39_sans_fonds_propres_declare_tout_et_le_signale(classeur):
    """Sans exercice antérieur, le seuil n'est pas mesurable."""

    groupes, parties = _groupes_parties_liees()
    reserves = _remplir_ep39(classeur, groupes, parties, 0.0)

    assert any("seuil de 5 %" in reserve.message for reserve in reserves)
    feuille = classeur["EP39"]
    lignes = indexer_codes_dispru(feuille)
    codes = sorted(code for code in lignes if code.startswith("PR"))
    noms = {feuille.cell(row=lignes[code], column=2).value for code in codes[:-1]}
    assert {"M. DIALLO Fatou", "Mme TRAORÉ Ali"} <= noms


def test_le_seuil_de_l_ep39_est_celui_du_formulaire():
    assert SEUIL_EP39 == 0.05


# ─── EP33 : le levier compte les derives ──────────────────────────────────


def test_le_levier_compte_l_exposition_des_derives(classeur):
    """Un dérivé déclaré sur l'EP11 ne peut pas valoir zéro sur l'EP33."""

    synthese = agregation.SyntheseCredit()
    ratio = _remplir_ep33(classeur, synthese, 10_000_000_000.0, 2_000_000_000.0)

    feuille = classeur["EP33"]
    lignes = indexer_codes_dispru(feuille)
    assert feuille.cell(row=lignes["RL005"], column=COLONNE_C).value == 2_000
    assert feuille.cell(row=lignes["RL007"], column=COLONNE_C).value == 2_000
    # L'exposition totale du levier les porte, et le ratio s'en trouve abaissé.
    assert feuille.cell(row=lignes["RL015"], column=COLONNE_C).value == 2_000
    assert ratio == pytest.approx(10_000_000_000.0 / 2_000_000_000.0)


# ─── EP12 a EP16 : la ponderation des derives ─────────────────────────────


def test_une_ponderation_absente_du_formulaire_monte_a_la_ligne_superieure():
    """L'état n'imprime que cinq ou six pondérations.

    Une exposition pondérée à 35 % sur un état qui ne propose ni 35 ni 40 ne
    doit pas disparaître : elle monte à 50 %, jamais elle ne descend.
    """

    bloc = _bloc_derives({0.35: 1_000.0}, (0.0, 0.2, 0.5, 1.0, 1.5))
    assert dict(bloc.ventilation.avant_arc) == {0.5: 1_000.0}


def test_sans_derive_le_bloc_reste_absent():
    assert _bloc_derives(None, (0.0, 0.2, 1.0)) is None
    assert _bloc_derives({}, (0.0, 0.2, 1.0)) is None


# ─── EP21 : le registre du produit brut face au compte de resultat ────────


def _exercice(annee: int, produit_brut: float):
    """Un exercice du registre de l'approche indicateur de base."""

    return SimpleNamespace(annee=annee, pnb_retenu_aib=produit_brut)


def test_un_produit_brut_qui_ne_ressemble_pas_au_compte_de_resultat_se_signale():
    """Un registre resté à des montants d'essai ne doit pas passer inaperçu.

    L'APR opérationnel suit le produit brut saisi : sur le portefeuille de
    référence, un registre à 456 789 FCFA déclarait 3 M d'actifs pondérés pour
    un compte de résultat portant 84 Md de PNB — soit trois millièmes de pour
    cent de l'APR total, sans que rien ne le dise.
    """

    with mock.patch(
        "app.rapports.fodep.service._pnb_des_etats_financiers",
        return_value={2024: 78_000_000_000.0, 2025: 84_000_000_000.0},
    ):
        reserves = _ecarts_de_produit_brut(
            [_exercice(2024, 456_789.0), _exercice(2025, 84_000_000_000.0)]
        )

    assert len(reserves) == 1
    message = reserves[0].message
    # 2025 concorde au franc pres : seul 2024 doit etre nomme.
    assert "2024" in message
    assert "2025" not in message
    assert reserves[0].nature is NatureReserve.A_VERIFIER


def test_un_ecart_de_retraitement_ne_se_signale_pas():
    """L'article 301 retranche et ajoute : quelques pour cent d'écart sont normaux."""

    with mock.patch(
        "app.rapports.fodep.service._pnb_des_etats_financiers",
        return_value={2025: 84_000_000_000.0},
    ):
        assert _ecarts_de_produit_brut([_exercice(2025, 80_000_000_000.0)]) == []


def test_sans_etats_financiers_rien_ne_se_compare():
    with mock.patch(
        "app.rapports.fodep.service._pnb_des_etats_financiers", return_value={}
    ):
        assert _ecarts_de_produit_brut([_exercice(2025, 1.0)]) == []
