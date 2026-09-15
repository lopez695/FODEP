"""Le registre des derives, et l'etat EP11 qu'il alimente.

L'EP11 declare le risque de contrepartie porte par les derives. Il etait saisi
cellule par cellule, puis plus du tout : la saisie manuelle a ete retiree de
l'ecran, et l'etat partait a zero sans qu'on puisse le corriger. Le formulaire
affirmait ainsi que l'etablissement ne detient aucun derive.

Le registre remplace la saisie de cellules par la saisie de contrats. Ce qui se
verifie ici, c'est le trajet d'un contrat jusqu'a sa case : la bonne ligne, la
bonne colonne, la ponderation du formulaire et pas une autre.
"""

from __future__ import annotations

from datetime import date, timedelta
from functools import cache
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.derives.models import DeriveCreate, DeriveUpdate, tranche_de_duree
from app.derives.services import (
    agreger_pour_ep11,
    categorie_ep11,
    lister_contreparties_derives,
    lister_sous_jacents,
    creer_derive,
    lister_derives,
    modifier_derive,
    supprimer_derive,
)
from app.rapports.fodep import construire_fodep
from app.rapports.fodep.disposition import indexer_codes_dispru
from app.rapports.fodep.service import synthese_ep11
from database.connection import database_manager


@pytest.fixture(scope="module", autouse=True)
def base_migree():
    """Applique les migrations : la table des derives est recente."""

    database_manager.initialize()


@pytest.fixture
def registre_rendu():
    """Rend au registre l'etat ou l'essai l'a trouve."""

    crees: list[int] = []
    yield crees
    for identifiant in crees:
        try:
            supprimer_derive(identifiant)
        except Exception:  # noqa: BLE001 - le test a pu le supprimer lui-meme
            pass


@pytest.fixture
def registre_mis_de_cote():
    """Vide le registre le temps de l'essai, puis rend les contrats a l'identique.

    La base d'essai est celle de l'application : les contrats qu'on y a saisis
    a la main ne doivent ni faire echouer l'essai, ni disparaitre apres lui.
    """

    with database_manager.read_connection() as connexion:
        curseur = connexion.execute("SELECT * FROM derives")
        colonnes = [description[0] for description in curseur.description]
        lignes = [tuple(ligne) for ligne in curseur.fetchall()]
    with database_manager.transaction() as connexion:
        connexion.execute("DELETE FROM derives")
    try:
        yield
    finally:
        marques = ", ".join("?" for _ in colonnes)
        with database_manager.transaction() as connexion:
            connexion.execute("DELETE FROM derives")
            connexion.executemany(
                f"INSERT INTO derives ({', '.join(colonnes)}) VALUES ({marques})",
                lignes,
            )


@cache
def _contrepartie_de(lettre: str) -> str:
    """Une contrepartie du portefeuille dont la categorie tombe sur la colonne.

    Les essais prennent des tiers reels plutot que d'en inventer : un derive se
    rattache a une contrepartie existante, et c'est ce trajet qu'on verifie. La
    categorie est lue par la fonction du moteur de calcul, comme a l'export.
    """

    from database.services.rwa_calculation_service import resolve_category

    with database_manager.read_connection() as connexion:
        lignes = connexion.execute(
            "SELECT id, categorie_prudentielle FROM contreparties ORDER BY id"
        ).fetchall()
    for ligne in lignes:
        if resolve_category(ligne["categorie_prudentielle"] or "")["code"] == lettre:
            return ligne["id"]
    pytest.skip(f"aucune contrepartie de categorie {lettre} dans la base d'essai")


def _contrat(categorie: str = "d", **champs) -> DeriveCreate:
    defauts = dict(
        contrepartie_id=_contrepartie_de(categorie),
        nature="taux",
        type_contrat="Swap de taux",
        devise="XOF",
        montant_notionnel=10_000_000_000.0,
        cout_remplacement=1_000_000_000.0,
        date_echeance=date(2030, 6, 30),
    )
    defauts.update(champs)
    return DeriveCreate(**defauts)


# ─── La tranche de duree ─────────────────────────────────────────────────────


def test_la_tranche_se_lit_en_duree_residuelle_pas_en_duree_d_origine():
    """Un contrat vieillit, et change de ligne en vieillissant.

    L'EP11 range ses engagements par duree RESIDUELLE. Un contrat a sept ans
    conclu il y a trois ans se declare aujourd'hui sur « > 1 an jusqu'a 5 ans »,
    pas sur « > 5 ans ». C'est pourquoi le registre garde l'echeance et non une
    tranche : une tranche stockee aurait gele la declaration au jour de la
    saisie.
    """

    echeance = date(2030, 1, 1)
    assert tranche_de_duree(echeance, date(2029, 6, 1)) == "moins_1_an"
    assert tranche_de_duree(echeance, date(2026, 1, 1)) == "1_a_5_ans"
    assert tranche_de_duree(echeance, date(2020, 1, 1)) == "plus_5_ans"


def test_un_contrat_echu_ne_sort_pas_du_registre():
    """Il tombe dans la premiere tranche, il ne disparait pas.

    Le supprimer serait perdre une trace ; le ranger ailleurs serait le
    declarer faux.
    """

    assert tranche_de_duree(date(2020, 1, 1), date(2026, 1, 1)) == "moins_1_an"


def test_les_bornes_tombent_du_bon_cote():
    """« < 1 an » et « > 1 an jusqu'a 5 ans » se touchent : la limite compte."""

    reference = date(2026, 1, 1)
    assert tranche_de_duree(reference + timedelta(days=364), reference) == "moins_1_an"
    assert tranche_de_duree(reference + timedelta(days=400), reference) == "1_a_5_ans"
    assert tranche_de_duree(reference + timedelta(days=1820), reference) == "1_a_5_ans"
    assert tranche_de_duree(reference + timedelta(days=1900), reference) == "plus_5_ans"


# ─── L'agregation ────────────────────────────────────────────────────────────


def test_chaque_contrat_rejoint_la_case_de_sa_nature_et_de_sa_duree(
    registre_rendu, registre_mis_de_cote
):
    """Cinq natures, trois durees : quinze cases, et une seule par contrat."""

    reference = date(2026, 6, 30)
    contrats = [
        _contrat(nature="taux", date_echeance=date(2027, 1, 1)),
        _contrat(nature="change_or", date_echeance=date(2027, 1, 1)),
        _contrat(nature="titres_propriete", date_echeance=date(2040, 1, 1)),
    ]
    for contrat in contrats:
        registre_rendu.append(creer_derive(contrat).id)

    agregats = agreger_pour_ep11(reference)
    assert agregats[("taux", "moins_1_an")].nombre_contrats == 1
    assert agregats[("change_or", "moins_1_an")].nombre_contrats == 1
    assert agregats[("titres_propriete", "plus_5_ans")].nombre_contrats == 1
    assert ("taux", "plus_5_ans") not in agregats


def test_deux_contrats_de_la_meme_case_s_additionnent(
    registre_rendu, registre_mis_de_cote
):
    """L'EP11 declare des lignes, pas des contrats."""

    reference = date(2026, 6, 30)
    for _ in range(2):
        registre_rendu.append(
            creer_derive(
                _contrat(
                    montant_notionnel=5e9,
                    cout_remplacement=1e9,
                    date_echeance=date(2027, 1, 1),
                )
            ).id
        )

    agregat = agreger_pour_ep11(reference)[("taux", "moins_1_an")]
    assert agregat.nombre_contrats == 2
    assert agregat.montant_notionnel == pytest.approx(10e9)
    assert agregat.cout_remplacement == pytest.approx(2e9)


# ─── L'etat produit ──────────────────────────────────────────────────────────


def test_la_ponderation_vient_du_formulaire_et_pas_du_code(registre_rendu):
    """« d = b x c » : c est imprime par la BCEAO, pas choisi par nous.

    RC051 -- instruments de taux, duree > 5 ans -- porte « 1,5 % » sur le
    formulaire. Coder ce taux ici en ferait une seconde source, qu'une nouvelle
    version du formulaire ferait diverger sans que rien ne le signale.
    """

    synthese = synthese_ep11(date(2026, 6, 30))
    par_code = {ligne.code: ligne for ligne in synthese.lignes}
    assert par_code["RC049"].ponderation == pytest.approx(0.0)
    assert par_code["RC050"].ponderation == pytest.approx(0.005)
    assert par_code["RC051"].ponderation == pytest.approx(0.015)
    assert par_code["RC052"].ponderation == pytest.approx(0.01)
    assert par_code["RC063"].ponderation == pytest.approx(0.15)


def test_l_exposition_est_le_cout_de_remplacement_plus_le_notionnel_pondere(
    registre_rendu, registre_mis_de_cote
):
    """e = a + d, et d = b x c. C'est la methode de l'exposition courante."""

    arrete = date(2026, 6, 30)
    registre_rendu.append(
        creer_derive(
            _contrat(
                nature="taux",
                # Echeance a plus de cinq ans : RC051, pondere a 1,5 %.
                date_echeance=date(2040, 1, 1),
                montant_notionnel=10_000_000_000.0,
                cout_remplacement=1_000_000_000.0,
            )
        ).id
    )

    ligne = {l.code: l for l in synthese_ep11(arrete).lignes}["RC051"]
    assert ligne.notionnel_pondere == pytest.approx(150_000_000.0)
    assert ligne.exposition == pytest.approx(1_150_000_000.0)


def test_la_ventilation_par_contrepartie_retrouve_toujours_l_exposition(
    registre_rendu,
):
    """L'etat ne peut plus se contredire d'une colonne a l'autre.

    Quand ces colonnes se saisissaient a la main, l'export devait verifier
    qu'elles bouclaient et signaler quand ce n'etait pas le cas. Elles se
    deduisent desormais des contrats : la verification devient un invariant.
    """

    arrete = date(2026, 6, 30)
    for categorie, notionnel in (("a", 6e9), ("d", 3e9), ("e", 1e9)):
        registre_rendu.append(
            creer_derive(
                _contrat(
                    categorie=categorie,
                    nature="change_or",
                    date_echeance=date(2027, 1, 1),
                    montant_notionnel=notionnel,
                    cout_remplacement=notionnel / 10,
                )
            ).id
        )

    ligne = {l.code: l for l in synthese_ep11(arrete).lignes}["RC052"]
    assert sum(ligne.ventilation.values()) == pytest.approx(ligne.exposition)
    assert set(ligne.ventilation) == {"a", "d", "e"}


def test_le_registre_arrive_dans_la_declaration(
    registre_rendu, registre_mis_de_cote
):
    """Le trajet complet : un contrat saisi, une case du classeur remplie."""

    registre_rendu.append(
        creer_derive(
            _contrat(
                nature="taux",
                date_echeance=date(2040, 1, 1),
                montant_notionnel=10_000_000_000.0,
                cout_remplacement=1_000_000_000.0,
                categorie="d",
            )
        ).id
    )

    resultat = construire_fodep(date(2026, 6, 30))
    feuille = load_workbook(BytesIO(resultat.contenu))["EP11"]
    lignes = indexer_codes_dispru(feuille)

    def montant(code: str, colonne: int) -> float:
        return float(feuille.cell(row=lignes[code], column=colonne).value or 0)

    # Les montants sont declares en millions de FCFA (notice, § 2.3).
    assert montant("RC051", 3) == pytest.approx(1_000, abs=1)  # (a)
    assert montant("RC051", 4) == pytest.approx(10_000, abs=1)  # (b)
    assert montant("RC051", 6) == pytest.approx(150, abs=1)  # (d) = b x 1,5 %
    assert montant("RC051", 7) == pytest.approx(1_150, abs=1)  # (e) = a + d
    # Institutions financieres : la colonne K.
    assert montant("RC051", 11) == pytest.approx(1_150, abs=1)

    # Et la ligne de total somme les quinze lignes, colonne par colonne.
    assert montant("RC064", 7) >= montant("RC051", 7)


def test_un_registre_vide_declare_zero_et_le_dit(registre_mis_de_cote):
    """Un zero transmis affirme quelque chose : il doit se lire.

    C'est la difference avec l'etat d'avant, ou l'EP11 partait a zero faute de
    pouvoir faire autrement.
    """

    assert not lister_derives(), "le registre doit etre vide pour cet essai"
    synthese = synthese_ep11(date(2026, 6, 30))
    assert synthese.nombre == 0
    assert synthese.total_exposition == 0
    assert any("registre est vide" in alerte for alerte in synthese.alertes)

    resultat = construire_fodep(date(2026, 6, 30))
    annonces = [
        reserve.message
        for reserve in resultat.anomalies
        if "registre des dérivés est vide" in reserve.message
    ]
    assert annonces, "un EP11 vide doit se dire"

    # Et l'etat ne figure plus parmi ceux qu'on declare a zero faute de source :
    # il en a une, elle est vide.
    sans_source = [
        reserve.message
        for reserve in resultat.anomalies
        if "déclarés à zéro faute de source" in reserve.message
    ]
    assert sans_source and "EP11" not in sans_source[0]


# ─── Le registre lui-meme ────────────────────────────────────────────────────


def test_un_contrat_se_cree_se_modifie_et_se_supprime(registre_rendu):
    identifiant = _contrepartie_de("d")
    contrat = creer_derive(_contrat("d"))
    registre_rendu.append(contrat.id)

    # Le nom vient de la fiche de la contrepartie, pas d'une saisie.
    with database_manager.read_connection() as connexion:
        nom = connexion.execute(
            "SELECT nom FROM contreparties WHERE id = ?", (identifiant,)
        ).fetchone()["nom"]
    assert contrat.contrepartie_id == identifiant
    assert contrat.contrepartie == nom

    modifie = modifier_derive(
        contrat.id, DeriveUpdate(montant_notionnel=42e9)
    )
    assert modifie.montant_notionnel == pytest.approx(42e9)
    assert modifie.contrepartie_id == identifiant, (
        "une modification partielle ne touche que ses champs"
    )

    supprimer_derive(contrat.id)
    registre_rendu.remove(contrat.id)
    assert all(ligne.id != contrat.id for ligne in lister_derives())


def test_un_contrat_en_devise_est_ramene_au_franc_cfa(registre_rendu):
    """Le formulaire se declare en FCFA, quelle que soit la devise du contrat."""

    reference = date(2026, 6, 30)
    registre_rendu.append(
        creer_derive(
            _contrat(
                devise="EUR",
                montant_notionnel=1_000_000.0,
                cout_remplacement=0.0,
                date_echeance=date(2027, 1, 1),
            )
        ).id
    )
    agregat = agreger_pour_ep11(reference)[("taux", "moins_1_an")]
    assert agregat.montant_notionnel > 1_000_000.0, (
        "un notionnel en euro doit valoir davantage une fois converti en FCFA"
    )


# ─── Le rattachement aux contreparties ───────────────────────────────────────


def test_la_colonne_vient_de_la_categorie_de_la_contrepartie(registre_rendu):
    """Un derive sur un souverain se ventile chez les souverains.

    La categorie ne se saisit plus : elle se lit sur la fiche, par la fonction
    qu'emploie le moteur de calcul. Un derive et un pret sur le meme tiers ne
    peuvent plus tomber dans deux categories differentes.
    """

    contrat = creer_derive(_contrat("a"))
    registre_rendu.append(contrat.id)
    assert contrat.categorie_contrepartie == "a"
    assert not contrat.hors_ep11


def test_une_contrepartie_inconnue_est_refusee():
    """Un derive sur un tiers que l'application ne connait pas ne se relie a rien."""

    with pytest.raises(HTTPException) as erreur:
        creer_derive(_contrat(contrepartie_id="EXP-INCONNUE-0000"))
    assert erreur.value.status_code == 422
    assert erreur.value.detail["code"] == "CONTREPARTIE_INCONNUE"


def test_changer_de_contrepartie_change_la_colonne(registre_rendu):
    contrat = creer_derive(_contrat("d"))
    registre_rendu.append(contrat.id)
    assert contrat.categorie_contrepartie == "d"

    modifie = modifier_derive(
        contrat.id, DeriveUpdate(contrepartie_id=_contrepartie_de("a"))
    )
    assert modifie.categorie_contrepartie == "a"


def test_une_categorie_sans_colonne_se_range_sous_entreprises(registre_rendu):
    """La clientele de detail n'a pas de colonne sur l'EP11.

    Le derive se range sous « Entreprises », repli du formulaire -- mais ce
    n'est pas la categorie du tiers, et l'ecran comme l'export le disent.
    """

    contrat = creer_derive(_contrat("f", date_echeance=date(2027, 1, 1)))
    registre_rendu.append(contrat.id)
    assert contrat.categorie_contrepartie == "e"
    assert contrat.hors_ep11

    synthese = synthese_ep11(date(2026, 6, 30))
    assert any("pas de colonne" in alerte for alerte in synthese.alertes)


def test_le_selecteur_propose_les_contreparties_avec_leur_identifiant():
    """C'est par l'identifiant qu'on fait la correspondance avec le reste."""

    contreparties = lister_contreparties_derives()
    assert contreparties, "le portefeuille d'essai porte des contreparties"
    assert all(c.id for c in contreparties)
    assert all(c.categorie_ep11 in {"a", "b", "c", "d", "e"} for c in contreparties)
    # Et la colonne proposee est celle que prendra un contrat enregistre.
    premiere = contreparties[0]
    assert categorie_ep11(premiere.categorie_prudentielle)[0] == premiere.categorie_ep11


def test_une_devise_que_l_outil_ne_sait_pas_convertir_est_refusee():
    """Une devise inconnue se convertirait au taux de repli 1,0.

    Un contrat en livres partirait alors dans l'EP11 comme s'il etait libelle en
    francs. L'ecran ne propose que les devises convertibles ; le serveur s'y
    tient aussi, pour un appel direct a l'API.
    """

    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        _contrat(devise="GBP")
    with pytest.raises(ValidationError):
        DeriveUpdate(devise="GBP")
    assert _contrat(devise="eur").devise == "EUR"


# ─── Ce que le derive couvre ─────────────────────────────────────────────────


def test_le_selecteur_propose_credits_obligations_et_actions():
    """Les trois onglets du formulaire, avec leurs identifiants.

    Les obligations et les actions sont celles du portefeuille de marche, lues
    par le module VaR -- pas une table a part qui pourrait etre perimee.
    """

    sous_jacents = lister_sous_jacents()
    assert sous_jacents.credits
    assert all(c.id and c.contrepartie_id for c in sous_jacents.credits)
    assert all(c.categorie_ep11 in {"a", "b", "c", "d", "e"} for c in sous_jacents.credits)
    assert all(o.isin and o.emetteur for o in sous_jacents.obligations)
    assert all(a.ticker and a.libelle for a in sous_jacents.actions)


def test_un_credit_se_couvre_avec_son_propre_client(registre_rendu):
    """Le client qui emprunte signe la couverture : c'est lui la contrepartie."""

    credits = lister_sous_jacents().credits
    credit = credits[0]
    contrat = creer_derive(_contrat(
        contrepartie_id=credit.contrepartie_id,
        sous_jacent_type="credit",
        sous_jacent_ref=credit.id,
    ))
    registre_rendu.append(contrat.id)
    assert contrat.sous_jacent_type == "credit"
    assert contrat.sous_jacent_ref == credit.id
    assert credit.id in (contrat.sous_jacent_libelle or "")

    autre = next(c for c in credits if c.contrepartie_id != credit.contrepartie_id)
    with pytest.raises(HTTPException) as erreur:
        creer_derive(_contrat(
            contrepartie_id=autre.contrepartie_id,
            sous_jacent_type="credit",
            sous_jacent_ref=credit.id,
        ))
    assert erreur.value.status_code == 422
    assert erreur.value.detail["code"] == "CONTREPARTIE_INCOHERENTE"


def test_un_derive_sur_obligation_porte_sur_les_taux(registre_rendu):
    """Et sa contrepartie reste celui qui a signe, pas l'emetteur.

    Un swap qui couvre une obligation d'Etat est signe avec une banque : le
    ventiler chez les souverains le ponderait plus tard a 0 %.
    """

    obligations = lister_sous_jacents().obligations
    if not obligations:
        pytest.skip("aucune obligation dans le portefeuille d'essai")
    obligation = obligations[0]

    contrat = creer_derive(_contrat(
        "d",
        sous_jacent_type="obligation",
        sous_jacent_ref=obligation.isin,
        nature="taux",
    ))
    registre_rendu.append(contrat.id)
    assert contrat.sous_jacent_ref == obligation.isin
    assert contrat.categorie_contrepartie == "d"

    with pytest.raises(HTTPException) as erreur:
        creer_derive(_contrat(
            sous_jacent_type="obligation",
            sous_jacent_ref=obligation.isin,
            nature="change_or",
        ))
    assert erreur.value.detail["code"] == "NATURE_INCOMPATIBLE"

    with pytest.raises(HTTPException) as erreur:
        creer_derive(_contrat(
            sous_jacent_type="obligation", sous_jacent_ref="ISIN-INCONNU"
        ))
    assert erreur.value.detail["code"] == "SOUS_JACENT_INCONNU"


def test_un_derive_sur_action_porte_sur_les_titres(registre_rendu):
    actions = lister_sous_jacents().actions
    if not actions:
        pytest.skip("aucune action dans le portefeuille d'essai")
    action = actions[0]

    contrat = creer_derive(_contrat(
        sous_jacent_type="action",
        sous_jacent_ref=action.ticker,
        nature="titres_propriete",
    ))
    registre_rendu.append(contrat.id)
    assert contrat.sous_jacent_libelle == action.libelle

    with pytest.raises(HTTPException) as erreur:
        creer_derive(_contrat(
            sous_jacent_type="action", sous_jacent_ref=action.ticker, nature="taux"
        ))
    assert erreur.value.detail["code"] == "NATURE_INCOMPATIBLE"


def test_changer_de_nature_reverifie_le_sous_jacent(registre_rendu):
    """Une modification ne doit pas rendre le contrat incoherent en douce."""

    obligations = lister_sous_jacents().obligations
    if not obligations:
        pytest.skip("aucune obligation dans le portefeuille d'essai")
    contrat = creer_derive(_contrat(
        sous_jacent_type="obligation",
        sous_jacent_ref=obligations[0].isin,
        nature="taux",
    ))
    registre_rendu.append(contrat.id)

    with pytest.raises(HTTPException) as erreur:
        modifier_derive(contrat.id, DeriveUpdate(nature="change_or"))
    assert erreur.value.detail["code"] == "NATURE_INCOMPATIBLE"
