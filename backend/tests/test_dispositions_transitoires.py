"""Les dispositions transitoires sur les fonds propres, et l'etat EP04.

Bale III a rendu inadmissibles certains elements de fonds propres au
1er janvier 2018 et les retire par paliers : un taux, applique au stock de
l'epoque, dit combien on peut encore en compter, et le montant encore en
circulation le plafonne.

L'application n'en tenait aucune trace. Trois cles de l'EP03 --
`fpi07_transitoire`, `fpi25_transitoire`, `fpi33_transitoire` -- etaient lues
sur la table des fonds propres, qui ne les a jamais portees : elles valaient
zero a chaque export, et l'EP04 partait a zero avec elles. Ce qui se verifie
ici, c'est le trajet d'un montant saisi jusqu'aux fonds propres declares.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.dispositions_transitoires.models import (
    DispositionsTransitoiresUpdate,
    DispositionsTransitoiresView,
)
from app.dispositions_transitoires.services import (
    calculer_ep04,
    enregistrer_dispositions,
    lire_dispositions,
    supprimer_dispositions,
)
from app.rapports.fodep import construire_fodep
from app.rapports.fodep.disposition import CHEMIN_MODELE, indexer_codes_dispru
from app.rapports.fodep.service import _taux_de_retrait, synthese_ep04
from database.connection import database_manager

EXERCICE_ESSAI = 2099

#: L'export lit les dispositions de l'exercice DECLARE : l'arrete de l'essai
#: doit donc tomber dans le millesime sur lequel il a saisi.
ARRETE_ESSAI = date(EXERCICE_ESSAI, 12, 31)


@pytest.fixture(scope="module", autouse=True)
def base_migree():
    """Applique les migrations : la table des dispositions est recente."""

    database_manager.initialize()


@pytest.fixture
def dispositions_rendues():
    """Rend a la table l'etat ou l'essai l'a trouvee."""

    exercices: list[int] = []
    yield exercices
    for exercice in exercices:
        supprimer_dispositions(exercice)


@pytest.fixture
def table_videe():
    """Vide la table le temps de l'essai, puis la remet comme elle etait.

    La base des tests est recopiee depuis la graine versionnee, et celle-ci
    porte les saisies faites depuis l'ecran. Un essai qui affirme « le registre
    est vide » sans le rendre vide ne teste donc pas ce qu'il croit : il passe
    tant que personne n'a rien saisi, et echoue le jour ou quelqu'un s'en sert.
    """

    with database_manager.read_connection() as connexion:
        conservees = [
            dict(ligne)
            for ligne in connexion.execute(
                "SELECT * FROM dispositions_transitoires"
            ).fetchall()
        ]
    with database_manager.transaction() as connexion:
        connexion.execute("DELETE FROM dispositions_transitoires")

    yield

    with database_manager.transaction() as connexion:
        connexion.execute("DELETE FROM dispositions_transitoires")
        for ligne in conservees:
            colonnes = ", ".join(ligne)
            valeurs = ", ".join("?" * len(ligne))
            connexion.execute(
                f"INSERT INTO dispositions_transitoires({colonnes}) "
                f"VALUES ({valeurs})",
                list(ligne.values()),
            )


def _saisie(**champs) -> DispositionsTransitoiresUpdate:
    defauts = dict(
        exercice=EXERCICE_ESSAI,
        part_capital_non_admissible=10_000_000_000.0,
        provisions_reglementees=2_000_000_000.0,
        fonds_affectes=1_000_000_000.0,
        cet1_en_circulation=20_000_000_000.0,
        cet1_eligible_at1=3_000_000_000.0,
        cet1_eligible_t2_autres=1_000_000_000.0,
        cet1_exclu=500_000_000.0,
        dettes_subordonnees_2018=8_000_000_000.0,
        part_dettes_non_admissible=4_000_000_000.0,
        ecarts_reevaluation=1_000_000_000.0,
        autres_t2_non_admissibles=0.0,
        t2_en_circulation=3_000_000_000.0,
    )
    defauts.update(champs)
    return DispositionsTransitoiresUpdate(**defauts)


# ─── Le taux de retrait ──────────────────────────────────────────────────────


def test_le_taux_de_retrait_vient_du_formulaire_et_pas_du_code():
    """Le taux (a) est imprime par la BCEAO sur la ligne DT001.

    Sa cellule est verrouillee : ce n'est ni au declarant ni a l'application de
    le choisir. Le retrait progressif s'abaisse d'un exercice a l'autre, et
    c'est la nouvelle version du formulaire qui le dira -- le coder ici en
    ferait une seconde source a corriger a chaque fois.
    """

    classeur = load_workbook(CHEMIN_MODELE, read_only=True)
    try:
        assert _taux_de_retrait(classeur) == pytest.approx(0.9)
    finally:
        classeur.close()


def test_la_ligne_du_taux_n_est_pas_ouverte_a_la_saisie():
    """Si la BCEAO l'ouvrait un jour, il faudrait le savoir."""

    from openpyxl.cell.cell import MergedCell

    from app.rapports.fodep.disposition import styles_ouverts

    classeur = load_workbook(CHEMIN_MODELE)
    try:
        styles = styles_ouverts(classeur)
        feuille = classeur["EP04"]
        rang = indexer_codes_dispru(feuille)["DT001"]
        cellule = feuille.cell(row=rang, column=3)
        assert not isinstance(cellule, MergedCell)
        assert cellule._style is not None
        assert cellule._style.protectionId not in (styles or frozenset()), (
            "DT001 est verrouillee : le taux de retrait fait foi tel quel."
        )
    finally:
        classeur.close()


# ─── Les formules de l'etat ──────────────────────────────────────────────────


def _vue(saisie: DispositionsTransitoiresUpdate) -> DispositionsTransitoiresView:
    """La saisie vue comme un enregistrement, sans passer par la base.

    Les formules se verifient sur des nombres, pas sur un aller-retour SQL.
    """

    return DispositionsTransitoiresView(
        **saisie.model_dump(), cree_le="", modifie_le=""
    )


def test_les_formules_sont_celles_que_le_formulaire_imprime():
    """« f = c+d+e », « g = f x a », « i = min(g,h) », et leurs jumelles en T2."""

    calcul = calculer_ep04(_vue(_saisie()), 0.9)
    assert calcul.total_cet1_non_admissible == pytest.approx(13e9)  # 10 + 2 + 1
    assert calcul.plafond_cet1 == pytest.approx(11.7e9)  # 13 x 0,9
    # 20 Md circulent encore, mais le taux n'en autorise que 11,7.
    assert calcul.cet1_reconnu == pytest.approx(11.7e9)

    assert calcul.total_t2_non_admissible == pytest.approx(5e9)  # 4 + 1 + 0
    assert calcul.plafond_t2 == pytest.approx(4.5e9)  # 5 x 0,9
    # Ici c'est l'inverse : le plafond autorise 4,5 Md, mais 3 seulement
    # circulent encore.
    assert calcul.t2_reconnu == pytest.approx(3e9)


def test_le_plafond_et_la_circulation_se_plafonnent_l_un_l_autre():
    """min(g, h) : ni plus que le taux n'autorise, ni plus qu'il n'en reste."""

    # Le taux mord : 13 Md x 0,9 = 11,7, contre 20 encore en circulation.
    mord_par_le_taux = calculer_ep04(_vue(_saisie()), 0.9)
    assert mord_par_le_taux.cet1_reconnu == pytest.approx(11.7e9)

    # La circulation mord : il ne reste que 5 Md des 13 d'origine.
    mord_par_la_circulation = calculer_ep04(
        _vue(_saisie(cet1_en_circulation=5e9)), 0.9
    )
    assert mord_par_la_circulation.cet1_reconnu == pytest.approx(5e9)


def test_sans_saisie_tout_vaut_zero():
    """Un etablissement sans instrument en retrait progressif declare zero."""

    calcul = calculer_ep04(None, 0.9)
    assert calcul.cet1_reconnu == 0.0
    assert calcul.t2_reconnu == 0.0
    assert calcul.total_cet1_non_admissible == 0.0
    assert calcul.total_t2_non_admissible == 0.0


# ─── Le registre ─────────────────────────────────────────────────────────────


def test_un_exercice_un_enregistrement(dispositions_rendues):
    """Deux corrections du meme millesime se remplacent, elles ne s'empilent pas."""

    dispositions_rendues.append(EXERCICE_ESSAI)
    enregistrer_dispositions(_saisie(part_capital_non_admissible=1e9))
    enregistrer_dispositions(_saisie(part_capital_non_admissible=7e9))

    lu = lire_dispositions(EXERCICE_ESSAI)
    assert lu is not None
    assert lu.part_capital_non_admissible == pytest.approx(7e9)


def test_rien_de_saisi_se_distingue_de_tout_a_zero():
    """`None` et un enregistrement nul ne se declarent pas de la meme facon."""

    assert lire_dispositions(2098) is None


# ─── Le trajet jusqu'aux fonds propres declares ──────────────────────────────


def test_l_ep04_alimente_les_quatre_lignes_de_l_ep03(
    table_videe, dispositions_rendues
):
    """FPI07 en CET1, FPI25 en AT1, FPI33 et FPI34 en T2.

    C'est le point du dispositif : ce que les dispositions transitoires
    laissent encore compter entre reellement dans les fonds propres.
    """

    dispositions_rendues.append(EXERCICE_ESSAI)
    enregistrer_dispositions(_saisie())

    resultat = construire_fodep(ARRETE_ESSAI)
    classeur = load_workbook(BytesIO(resultat.contenu))
    feuille = classeur["EP03"]
    lignes = indexer_codes_dispru(feuille)

    def montant(code: str) -> float:
        return float(feuille.cell(row=lignes[code], column=3).value or 0)

    # Les montants sont declares en millions de FCFA (notice, § 2.3).
    assert montant("FPI07") == pytest.approx(11_700, abs=1)
    assert montant("FPI25") == pytest.approx(3_000, abs=1)
    assert montant("FPI33") == pytest.approx(1_000, abs=1)
    assert montant("FPI34") == pytest.approx(3_000, abs=1)

    # Et l'EP04 declare les memes chiffres que l'EP03 en reprend : deux etats
    # qui se contrediraient feraient rejeter la declaration.
    ep04 = classeur["EP04"]
    lignes_ep04 = indexer_codes_dispru(ep04)
    for code in ("FPI07", "FPI25", "FPI33", "FPI34"):
        declare = float(ep04.cell(row=lignes_ep04[code], column=3).value or 0)
        assert declare == pytest.approx(montant(code), abs=1), code


def test_le_transitoire_entre_dans_le_total_du_tier_2(
    table_videe, dispositions_rendues
):
    """FPI33 et FPI34 ne sont pas decoratifs : ils grossissent le T2.

    La table est videe le temps de l'essai : sans cela, la mesure de reference
    porterait deja un transitoire, et l'ecart mesure ne serait pas celui qu'on
    croit.
    """

    def total_t2() -> float:
        classeur = load_workbook(BytesIO(construire_fodep(ARRETE_ESSAI).contenu))
        feuille = classeur["EP03"]
        lignes = indexer_codes_dispru(feuille)
        return float(feuille.cell(row=lignes["FPI39"], column=3).value or 0)

    sans = total_t2()
    dispositions_rendues.append(EXERCICE_ESSAI)
    enregistrer_dispositions(_saisie())
    avec = total_t2()

    # FPI33 (1 000 M) + FPI34 (3 000 M).
    assert avec - sans == pytest.approx(4_000, abs=1)


def test_l_etat_n_est_plus_declare_a_zero_faute_de_source(
    table_videe, dispositions_rendues
):
    """L'EP04 a desormais une source : il sort de la liste des etats muets."""

    dispositions_rendues.append(EXERCICE_ESSAI)
    enregistrer_dispositions(_saisie())

    resultat = construire_fodep(ARRETE_ESSAI)
    sans_source = [
        reserve.message
        for reserve in resultat.anomalies
        if "déclarés à zéro faute de source" in reserve.message
    ]
    assert sans_source and "EP04" not in sans_source[0]


def test_les_dispositions_d_un_autre_exercice_ne_sont_pas_declarees(
    table_videe, dispositions_rendues
):
    """Une saisie porte sur un millesime, pas sur la table.

    L'export lisait « les dispositions les plus recentes ». Une saisie portee
    sur un exercice posterieur entrait donc dans une declaration qui n'etait pas
    la sienne, et gonflait ses fonds propres sans que rien ne le dise. C'est
    arrive : une saisie d'essai sur 2099 s'est retrouvee dans l'arrete de 2026.
    """

    dispositions_rendues.append(EXERCICE_ESSAI)
    enregistrer_dispositions(_saisie())

    # On declare un AUTRE exercice que celui qui porte la saisie.
    autre = date(EXERCICE_ESSAI - 1, 12, 31)
    classeur = load_workbook(BytesIO(construire_fodep(autre).contenu))
    feuille = classeur["EP03"]
    lignes = indexer_codes_dispru(feuille)

    for code in ("FPI07", "FPI25", "FPI33", "FPI34"):
        montant = float(feuille.cell(row=lignes[code], column=3).value or 0)
        assert montant == 0, (
            f"{code} porte {montant} M alors que la saisie vise "
            f"{EXERCICE_ESSAI}, pas {autre.year}."
        )


def test_un_registre_vide_declare_zero_et_le_dit(table_videe):
    """Un zero transmis affirme quelque chose : il doit se lire."""

    assert lire_dispositions() is None
    synthese = synthese_ep04(EXERCICE_ESSAI)
    assert not synthese.renseigne
    assert synthese.taux_de_retrait == pytest.approx(0.9)
    assert any("partira à zéro" in alerte for alerte in synthese.alertes)

    annonces = [
        reserve.message
        for reserve in construire_fodep(ARRETE_ESSAI).anomalies
        if "aucune disposition transitoire" in reserve.message
    ]
    assert annonces, "un EP04 vide doit se dire"


def test_l_apercu_distingue_ce_qui_est_saisi_calcule_ou_reporte(
    table_videe,
    dispositions_rendues,
):
    """L'ecran doit dire d'ou vient chaque montant.

    Un declarant qui cherche a corriger (i) doit comprendre qu'il ne se saisit
    pas, et que ce sont (g) et (h) qu'il faut reprendre.
    """

    dispositions_rendues.append(EXERCICE_ESSAI)
    enregistrer_dispositions(_saisie())

    par_code = {
        ligne.code: ligne for ligne in synthese_ep04(EXERCICE_ESSAI).lignes
    }
    assert par_code["FPI01"].origine == "reporte"  # le capital vient de l'EP03
    assert par_code["DT002"].origine == "saisie"
    assert par_code["DT005"].origine == "calcule"
    assert par_code["FPI07"].origine == "calcule"
    assert par_code["FPI07"].formule == "(i) = min(g, h)"
