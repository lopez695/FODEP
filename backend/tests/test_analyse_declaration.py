"""Analyse d'une declaration FODEP au regard du dispositif prudentiel UMOA.

Le texte d'essai est celui qu'une extraction reelle produit : les mots y sont
colles, sans espaces, parce qu'un PDF ne porte pas de separateurs de mots. C'est
pour cela que les intitules des normes sont pris dans le formulaire de la BCEAO
et non dans le PDF.

Le sens de chaque norme est ce qui decide de la conformite : un ratio de fonds
propres de 12 % est conforme face a un minimum de 11,50 %, une limite de
participation a 30 % ne l'est pas face a un plafond de 25 %. Ces tests verifient
donc les deux sens, pas seulement la lecture des nombres.
"""

from __future__ import annotations

import pytest

from app.rapports.fodep.analyse import (
    ARRONDI,
    DEPASSEE,
    ECART,
    MAXIMUMS,
    MINIMUMS,
    NON_MESUREE,
    RESPECTEE,
    analyser_classeur,
    analyser_declaration,
    analyser_texte,
    libelles_des_normes,
    sens_des_normes,
)

# Extrait tel que pypdf le rend sur une declaration produite par l'outil.
TEXTE_EP01 = (
    "EP01 4/59ETAT DE CONFORMITE AUX NORMES PRUDENTIELLES EP01Code DISPRU "
    "Liste des normes prudentielles Reference Niveau a respecter Niveau observe "
    "Situation de l'etablissementA. Normes de solvabilite"
    "RA001 RatiodefondspropresCET1(%) EP02 0,0750 0,1230"
    "RA002 RatiodefondspropresdebaseT1(%) EP02 0,0850 0,1313"
    "RA003 Ratiodesolvabilitetotal(%) EP02 0,1150 0,1509"
    "B. Norme de division des risques"
    "RA004 Normededivisiondesrisques EP29 0,2500 0,1644"
    "C. Ratio de levierRA005 Ratiodelevier EP33 0,0300 0,1297"
    "D. Autres normes prudentielles"
    "RA006 Limiteindividuellesurlesparticipations(25%capitaldel'entreprise) "
    "EP35 0,2500 0,2250"
    "RA007 Limiteindividuelle(15%desfondspropresT1) EP35 0,1500 0,0911"
    "RA008 Limiteglobale(60%desfondspropreseffectifs) EP35 0,6000 0,0792"
    "RA009 Limitesurlesimmobilisationshorsexploitation EP36 0,1500 0"
    "RA010 Limitesurletotaldesimmobilisationsetdesparticipations EP37 1 0,1062"
    "RA011 Limitesurlespretsauxactionnaires EP38 0,2000 0,0008"
)


@pytest.fixture(scope="module")
def normes():
    return {norme.code: norme for norme in analyser_texte(TEXTE_EP01)}


def test_les_onze_normes_sont_lues(normes):
    assert sorted(normes) == [f"RA{numero:03d}" for numero in range(1, 12)]


def test_les_seuils_et_les_niveaux_observes_sont_lus(normes):
    assert normes["RA001"].seuil == pytest.approx(0.075)
    assert normes["RA001"].observe == pytest.approx(0.1230)
    assert normes["RA001"].reference == "EP02"
    # Un seuil entier — 100 % pour RA010 — se lit aussi.
    assert normes["RA010"].seuil == pytest.approx(1.0)
    assert normes["RA010"].observe == pytest.approx(0.1062)


def test_les_pourcentages_de_l_intitule_ne_sont_pas_pris_pour_des_niveaux(normes):
    """« 25 % du capital » figure dans le libelle de RA006.

    Confondre ce rappel avec un niveau observe ferait declarer la norme sur son
    propre plafond, donc toujours conforme.
    """

    assert normes["RA006"].seuil == pytest.approx(0.25)
    assert normes["RA006"].observe == pytest.approx(0.2250)
    assert normes["RA007"].observe == pytest.approx(0.0911)


def test_le_sens_de_chaque_norme_decide_de_la_conformite(normes):
    for code in MINIMUMS:
        assert normes[code].minimum is True, code
    for code in MAXIMUMS:
        assert normes[code].minimum is False, code

    # 12,30 % pour un plancher de 7,50 % : conforme, avec 4,80 points de marge.
    assert normes["RA001"].situation == RESPECTEE
    assert normes["RA001"].ecart == pytest.approx(0.048)
    # 22,50 % pour un plafond de 25 % : conforme, 2,50 points de marge.
    assert normes["RA006"].situation == RESPECTEE
    assert normes["RA006"].ecart == pytest.approx(0.025)


def test_un_plancher_sous_son_seuil_est_depasse():
    texte = "RA003 Ratiodesolvabilitetotal(%) EP02 0,1150 0,0900"
    norme = analyser_texte(texte)[0]
    assert norme.situation == DEPASSEE
    # L'ecart est negatif : il manque 2,50 points pour atteindre le plancher.
    assert norme.ecart == pytest.approx(-0.025)


def test_un_plafond_au_dessus_de_son_seuil_est_depasse():
    texte = "RA006 Limiteindividuelle(25%capital) EP35 0,2500 0,3100"
    norme = analyser_texte(texte)[0]
    assert norme.situation == DEPASSEE
    assert norme.ecart == pytest.approx(-0.06)


def test_une_norme_sans_niveau_observe_est_non_mesuree():
    """La colonne « situation » est calculee par le formulaire : sa valeur
    n'existe qu'a l'ouverture du classeur. Sans niveau observe, une norme n'est
    ni conforme ni depassee — elle est inconnue."""

    norme = analyser_texte("RA004 Normededivisiondesrisques EP29 0,2500")[0]
    assert norme.observe is None
    assert norme.situation == NON_MESUREE
    assert norme.ecart is None


def test_les_intitules_viennent_du_formulaire(normes):
    """Le texte extrait arrive colle : le libelle exact vient du modele."""

    assert normes["RA001"].libelle == "Ratio de fonds propres CET 1 (%)"
    assert " " in normes["RA006"].libelle
    assert set(libelles_des_normes()) >= set(normes)


def test_le_sens_des_normes_est_lu_dans_le_formulaire():
    """La formule de la colonne « Situation » porte la regle de la BCEAO.

    `IF(G>F, "CONFORME", ...)` designe un plancher, `IF(G>F, "INFRACTION", ...)`
    un plafond. La lire evite de figer notre propre lecture du dispositif, et
    suit une revision du formulaire.
    """

    sens = sens_des_normes()
    assert len(sens) == 11
    assert {code for code, minimum in sens.items() if minimum} == set(MINIMUMS)
    assert {code for code, minimum in sens.items() if not minimum} == set(MAXIMUMS)


def test_a_egalite_exacte_un_plancher_suit_le_formulaire():
    """Le formulaire compare avec un « superieur » strict : un ratio pile au
    niveau a respecter y est declare INFRACTION. L'analyse doit dire la meme
    chose que la declaration, meme si « superieur ou egal » serait plus
    intuitif."""

    plancher = analyser_texte("RA001 Ratio EP02 0,1150 0,1150")[0]
    assert plancher.situation == DEPASSEE

    # Sur un plafond, le formulaire accepte l'egalite.
    plafond = analyser_texte("RA006 Limite EP35 0,2500 0,2500")[0]
    assert plafond.situation == RESPECTEE


# ─── Le classeur, piece transmise a la BCEAO ────────────────────────────────


@pytest.fixture(scope="module")
def classeur_reel() -> bytes:
    """Une declaration reelle, produite par l'outil lui-meme."""

    from app.rapports.fodep import construire_fodep

    return construire_fodep().contenu


def test_les_totaux_du_formulaire_sont_confrontes_a_leurs_lignes(classeur_reel):
    """Un total qui ne suit pas ses lignes fait rejeter le depot, sans qu'aucune
    norme ne soit en cause. Ces controles disent si la declaration se
    contredit."""

    analyse = analyser_declaration(classeur_reel, "FODEP.xlsx")
    assert analyse.classeur is True
    assert len(analyse.controles) >= 12

    # Aucun total ne doit s'ecarter au-dela de ce que les arrondis expliquent.
    ecarts = [c for c in analyse.controles if c.statut == ECART]
    assert not ecarts, [
        f"{c.etat} {c.colonne} : {c.constate} au lieu de {c.attendu}"
        for c in ecarts
    ]

    # Le formulaire arrondit chaque ligne au million : une derive d'une unite
    # sur dix-sept lignes est un arrondi, pas une erreur de somme. La tolerance
    # est arithmetique — une demi-unite par ligne sommee —, pas un chiffre rond
    # choisi pour faire passer le test.
    for controle in analyse.controles:
        assert controle.tolerance >= 1
        if controle.statut == ARRONDI:
            assert 0 < abs(controle.ecart) <= controle.tolerance


def test_un_total_fausse_est_signale(classeur_reel):
    """Le controle doit voir une somme qui ne tombe pas juste."""

    from io import BytesIO

    from openpyxl import load_workbook

    from app.rapports.fodep.disposition import indexer_codes_dispru

    classeur = load_workbook(BytesIO(classeur_reel))
    feuille = classeur["EP34"]
    lignes = indexer_codes_dispru(feuille)
    # Le total general des participations, gonfle de 700 : au-dela de tout
    # arrondi possible sur cinq lignes sommees.
    ligne = lignes["PA106"]
    actuel = feuille.cell(row=ligne, column=5).value or 0
    feuille.cell(row=ligne, column=5).value = actuel + 700
    tampon = BytesIO()
    classeur.save(tampon)

    analyse = analyser_declaration(tampon.getvalue(), "FODEP_fausse.xlsx")
    fautif = [
        controle
        for controle in analyse.controles
        if controle.etat == "EP34" and controle.statut == ECART
    ]
    assert fautif, "Le total gonfle n'a pas ete signale."
    assert fautif[0].ecart == pytest.approx(700)
    assert any("total" in message.lower() for message in analyse.avertissements)


def test_les_etats_entierement_a_zero_sont_recenses(classeur_reel):
    """Un etat a zero est declare comme les autres : rien n'y distingue un
    encours nul d'un encours non mesure. Le lecteur doit les voir."""

    analyse = analyser_declaration(classeur_reel, "FODEP.xlsx")
    assert len(analyse.inventaire) >= 30

    noms = {etat.nom for etat in analyse.inventaire}
    assert "EP01" in noms
    for etat in analyse.inventaire:
        assert etat.lignes > 0

    # La page de garde et l'attestation portent du texte, pas des zeros : les
    # signaler ferait passer une mise en page pour une declaration vide.
    textuelles = {"Page_de_garde", "ADPE"}
    for etat in analyse.inventaire:
        if etat.nom in textuelles:
            assert not etat.tout_a_zero, etat.nom


def test_une_impression_ne_porte_pas_de_controles():
    """Sommer des nombres extraits d'un PDF ne prouverait rien : on ne sait pas
    a quelle ligne du formulaire chacun appartient."""

    from io import BytesIO

    from pypdf import PdfWriter

    # Un PDF valide et vide, ecrit par pypdf lui-meme : un squelette bricole a
    # la main serait refuse au decodage, et le test porterait alors sur le
    # refus, pas sur l'absence de controles.
    ecrivain = PdfWriter()
    ecrivain.add_blank_page(width=595, height=842)
    tampon = BytesIO()
    ecrivain.write(tampon)

    analyse = analyser_declaration(tampon.getvalue(), "impression.pdf")
    assert analyse.classeur is False
    assert analyse.controles == []
    assert analyse.inventaire == []
    assert any("impression" in message for message in analyse.avertissements)


def test_le_classeur_se_lit_exactement(classeur_reel):
    """C'est la lecture de reference : des nombres, aux cases du formulaire.

    Aucune extraction de texte, donc aucune ambiguite — et c'est le classeur qui
    est transmis, pas son impression.
    """

    normes, controles, inventaire = analyser_classeur(classeur_reel)
    assert [norme.code for norme in normes] == [
        f"RA{numero:03d}" for numero in range(1, 12)
    ]
    assert controles, "Le classeur permet des controles de coherence."
    assert inventaire, "Le classeur permet d'inventorier ses etats."

    ra001 = normes[0]
    assert ra001.seuil == pytest.approx(0.075)
    assert ra001.observe is not None
    assert ra001.reference == "EP02"
    assert ra001.minimum is True


def _comme_une_impression(valeur: float | None) -> str:
    """Nombre tel qu'une impression du formulaire le porte."""

    return "" if valeur is None else f"{valeur:.4f}".replace(".", ",")


def test_les_deux_lectures_concluent_pareil(classeur_reel):
    """Le PDF est l'impression du classeur : sur les memes valeurs, les deux
    lectures doivent rendre le meme verdict, sinon l'une des deux se trompe.

    Le texte est reconstitue depuis le classeur lui-meme, et non copie d'une
    extraction figee : les donnees du poste evoluent, et comparer un echantillon
    d'hier a un classeur d'aujourd'hui ne prouverait rien.
    """

    depuis_classeur = analyser_declaration(classeur_reel, "FODEP.xlsx")
    assert depuis_classeur.pages == 0  # un classeur n'a pas de pages

    texte = "".join(
        f"{norme.code} Intitulecolle {norme.reference} "
        f"{_comme_une_impression(norme.seuil)} "
        f"{_comme_une_impression(norme.observe)}"
        for norme in depuis_classeur.normes
    )
    depuis_texte = {norme.code: norme for norme in analyser_texte(texte)}

    for norme in depuis_classeur.normes:
        lue = depuis_texte[norme.code]
        assert lue.seuil == pytest.approx(norme.seuil), norme.code
        if norme.observe is None:
            assert lue.observe is None, norme.code
        else:
            assert lue.observe == pytest.approx(norme.observe, abs=1e-4), norme.code
        assert lue.situation == norme.situation, norme.code
        assert lue.minimum == norme.minimum, norme.code


def test_un_classeur_sans_ep01_est_refuse():
    from io import BytesIO

    from openpyxl import Workbook

    classeur = Workbook()
    classeur.active.title = "Feuille1"
    tampon = BytesIO()
    classeur.save(tampon)

    with pytest.raises(ValueError, match="Aucun état EP01"):
        analyser_declaration(tampon.getvalue(), "autre.xlsx")


def test_un_fichier_qui_n_est_pas_une_declaration_est_refuse():
    with pytest.raises(ValueError, match="Format non reconnu"):
        analyser_declaration(b"bonjour", "note.txt")


def test_un_pdf_casse_le_dit():
    with pytest.raises(ValueError, match="PDF illisible"):
        analyser_declaration(b"%PDF-1.5 puis n'importe quoi", "casse.pdf")


def test_sans_bibliotheque_pdf_le_classeur_reste_analysable(
    monkeypatch, classeur_reel
):
    """Une dependance manquante ne doit pas emporter plus qu'elle ne sert.

    pypdf etait importe en tete de module : absent de l'environnement — le cas
    d'un venv ou il n'a pas ete installe —, il faisait echouer `app.main`, donc
    *toutes* les routes de l'API. Il n'est charge qu'a l'usage, et seule la
    lecture des PDF s'en trouve privee.
    """

    import sys

    monkeypatch.setitem(sys.modules, "pypdf", None)

    with pytest.raises(ValueError, match="pypdf n'est pas"):
        analyser_declaration(b"%PDF-1.5 quelque chose", "impression.pdf")

    # Le classeur, lui, se lit toujours : c'est la piece transmise.
    analyse = analyser_declaration(classeur_reel, "FODEP.xlsx")
    assert len(analyse.normes) == 11


def test_la_route_refuse_un_fichier_illisible_sans_planter():
    """Un depot errone doit se dire, pas faire tomber la requete en 500."""

    from fastapi.testclient import TestClient

    from app.main import app

    client = TestClient(app)
    reponse = client.post(
        "/rapports/fodep/analyse",
        files={"file": ("note.txt", b"bonjour", "text/plain")},
    )
    assert reponse.status_code == 422
    assert reponse.json()["detail"]["code"] == "FODEP_ANALYSE_PDF_ILLISIBLE"

    vide = client.post(
        "/rapports/fodep/analyse",
        files={"file": ("vide.pdf", b"", "application/pdf")},
    )
    assert vide.status_code == 422
    assert vide.json()["detail"]["code"] == "FODEP_ANALYSE_FICHIER_VIDE"
