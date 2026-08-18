"""L'export FODEP reproduit le formulaire de la BCEAO, renseigné et cohérent.

Le FODEP est transmis tel quel à la plate-forme de reporting : une colonne qui
ne se boucle pas ou un état dont le total ne correspond pas à ses lignes fait
rejeter la déclaration. Les contrôles ci-après portent donc sur les invariants
du formulaire lui-même, pas seulement sur le fait qu'un fichier soit produit.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO
import json
from unittest import mock

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Protection

from app.participations.models import ParticipationView
from app.rapports.fodep import construire_fodep, nom_fichier_fodep
from app.rapports.fodep.agregation import (
    PALIERS_EP20,
    repartir_sur_paliers,
)
from app.rapports.fodep.disposition import (
    completer_etat_a_zero,
    indexer_codes_dispru,
    lire_disposition_categorie,
    lire_etats_requis,
)
from app.rapports.fodep.service import (
    BASE_DE_DECLARATION,
    BORNES_ECHEANCE_EP31,
    CHEMIN_MODELE,
    CODES_EP22,
    COLONNES_PARTIES_LIEES,
    ETATS_ALIMENTES,
    ETATS_DECLARES_A_ZERO,
    PREMIERE_COLONNE_NUMERIQUE,
    SyntheseImmobilisations,
    _exercice_declare,
    _remplir_ep08,
    _remplir_ep22,
    _remplir_ep30,
    en_millions,
    repartir_en_millions,
)


ETATS_RISQUE_DE_CREDIT = ("EP12", "EP13", "EP14", "EP15", "EP16", "EP17", "EP18", "EP19")

COLONNE_AVANT_ARC = 3
COLONNE_GARANTIES = 4
COLONNE_APRES_ARC = 9
COLONNE_APR = 10

# Les montants sont déclarés en millions de FCFA : la somme d'un état et le
# total qu'il affiche peuvent différer d'un million par ligne arrondie.
TOLERANCE_ARRONDI = 10


@pytest.fixture(scope="module")
def classeur():
    resultat = construire_fodep()
    return load_workbook(BytesIO(resultat.contenu))


def test_le_modele_officiel_est_livre_avec_l_application():
    assert CHEMIN_MODELE.exists(), (
        "Le modèle FODEP doit accompagner le code : sans lui, l'export ne peut "
        "reproduire ni la mise en page ni les codes DISPRU du formulaire."
    )


def test_le_modele_ne_contient_aucune_donnee_declaree():
    """Le modèle provient d'un FODEP réel : il doit en avoir été purgé."""

    modele = load_workbook(CHEMIN_MODELE)
    residus: list[str] = []
    for nom in modele.sheetnames:
        if nom in {"ADPE", "EP01"}:
            # ADPE ne porte que des intitulés ; EP01 conserve les niveaux
            # réglementaires à respecter, qui font partie du formulaire.
            continue
        feuille = modele[nom]
        for ligne in feuille.iter_rows(min_row=8):
            for cellule in ligne:
                if cellule.value is None:
                    continue
                # Les notes de bas de page du formulaire — « * Y compris agios
                # dus », « *** de l'exercice précédent » — sont posées dans la
                # colonne des codes, que la BCEAO y laisse déverrouillée. Ce
                # sont des mentions imprimées, pas des montants déclarés : les
                # compter pour des résidus obligerait à les supprimer du
                # modèle, et l'export rendrait alors un formulaire amputé.
                if cellule.column == 1 and str(cellule.value).lstrip().startswith("*"):
                    continue
                if cellule.protection is not None and not cellule.protection.locked:
                    residus.append(f"{nom}!{cellule.coordinate}")
    assert not residus, f"Données résiduelles dans le modèle : {residus[:10]}"


def test_le_classeur_conserve_les_quarante_quatre_feuilles(classeur):
    assert len(classeur.sheetnames) == 44
    for etat in ("EP01", "EP02", "EP03", "EP08", "EP09", "EP10", "EP20", "EP21", "EP33"):
        assert etat in classeur.sheetnames


def test_la_colonne_garanties_se_boucle_a_zero(classeur):
    """Exigence de la notice : « la somme de chaque colonne doit être nulle ».

    Une garantie déplace l'exposition de la ligne du débiteur vers celle du
    garant. Si un transfert perd sa contrepartie — parce que la pondération du
    garant sort de la grille de l'état, par exemple — la colonne cesse de se
    boucler et l'état est incohérent.
    """

    for etat in ETATS_RISQUE_DE_CREDIT:
        feuille = classeur[etat]
        disposition = lire_disposition_categorie(feuille)
        somme = 0
        for bloc in disposition.blocs.values():
            for ligne in bloc.lignes_par_ponderation.values():
                somme += feuille.cell(row=ligne, column=COLONNE_GARANTIES).value or 0
        assert abs(somme) <= TOLERANCE_ARRONDI, (
            f"{etat} : la colonne « Garanties » totalise {somme} au lieu de 0."
        )


def test_les_actifs_ponderes_valent_l_exposition_fois_la_ponderation(classeur):
    for etat in ETATS_RISQUE_DE_CREDIT:
        feuille = classeur[etat]
        disposition = lire_disposition_categorie(feuille)
        for bloc in disposition.blocs.values():
            for ponderation, ligne in bloc.lignes_par_ponderation.items():
                apres_arc = feuille.cell(row=ligne, column=COLONNE_APRES_ARC).value or 0
                apr = feuille.cell(row=ligne, column=COLONNE_APR).value or 0
                assert abs(apr - apres_arc * ponderation) <= 1, (
                    f"{etat} ligne {ligne} : APR de {apr} pour une exposition de "
                    f"{apres_arc} pondérée à {ponderation:.0%}."
                )


def test_le_total_d_un_etat_egale_la_somme_de_ses_blocs(classeur):
    for etat in ETATS_RISQUE_DE_CREDIT:
        feuille = classeur[etat]
        disposition = lire_disposition_categorie(feuille)
        for colonne in (COLONNE_AVANT_ARC, COLONNE_APRES_ARC, COLONNE_APR):
            somme_blocs = sum(
                feuille.cell(row=bloc.ligne_total, column=colonne).value or 0
                for bloc in disposition.blocs.values()
            )
            total = (
                feuille.cell(
                    row=disposition.ligne_total_general, column=colonne
                ).value
                or 0
            )
            assert abs(total - somme_blocs) <= TOLERANCE_ARRONDI, (
                f"{etat} colonne {colonne} : total de {total} pour des blocs "
                f"totalisant {somme_blocs}."
            )


def test_l_ep08_totalise_les_etats_de_risque_de_credit(classeur):
    """Le total des APR doit être la somme de ses trois composantes."""

    feuille = classeur["EP08"]
    lignes = indexer_codes_dispru(feuille)

    def montant(code: str) -> float:
        return feuille.cell(row=lignes[code], column=5).value or 0

    somme_categories = sum(
        montant(code)
        for code in ("RC089", "RC114", "RC139", "RC164", "RC193", "RC219", "RC239", "RC261", "RC276")
    )
    assert abs(montant("APR01") - somme_categories) <= TOLERANCE_ARRONDI

    total_attendu = montant("APR01") + montant("APR02") + montant("APR03")
    assert abs(montant("APR04") - total_attendu) <= TOLERANCE_ARRONDI


CODES_DETAIL_MARCHE = ("RM044", "RM057", "RM067", "RM123")


def _remplir_ep08_avec(ventilation):
    """Renseigne un EP08 neuf et retourne ses montants de risque de marché."""

    classeur = load_workbook(CHEMIN_MODELE)
    _remplir_ep08(
        classeur,
        {"a": 1.0e11, "b": 2.0e10},
        5.0e9,
        194_688_063_683.7,
        3.0e10,
        ventilation,
    )
    feuille = classeur["EP08"]
    lignes = indexer_codes_dispru(feuille)
    montant = lambda code: feuille.cell(row=lignes[code], column=5).value or 0
    return [montant(code) for code in CODES_DETAIL_MARCHE], montant("APR02")


def test_le_detail_du_risque_de_marche_totalise_exactement_l_apr02():
    """Un état qui ne se boucle pas fait rejeter la déclaration.

    Les quatre natures sont déclarées en millions : arrondir chacune de son
    côté suffirait à décaler leur somme du total de quelques millions.
    """

    detail, total = _remplir_ep08_avec(
        {
            "taux": 1_120_456_789.0,
            "actions": 640_222_111.0,
            "change": 15_762_999_199.3,
            "produits_de_base": 0.0,
        }
    )
    assert all(montant > 0 for montant in detail[:3])
    assert sum(detail) == total, (
        "Les lignes RM044 à RM123 doivent totaliser l'APR02 au million près : "
        f"{sum(detail)} déclaré contre {total} attendu."
    )


def test_sans_ventilation_le_detail_du_risque_de_marche_reste_a_zero():
    """Faute de détail transmis, mieux vaut zéro qu'une répartition inventée."""

    detail, total = _remplir_ep08_avec(None)
    assert detail == [0, 0, 0, 0]
    assert total > 0, "Le total doit rester déclaré même sans son détail."


def test_la_repartition_en_millions_ne_perd_aucun_million():
    """Le reste de l'arrondi revient aux parts, il ne disparaît pas."""

    parts = {"taux": 1.0, "actions": 1.0, "change": 1.0, "produits_de_base": 0.0}
    repartition = repartir_en_millions(10_000_000_000.0, parts)
    assert sum(repartition.values()) == en_millions(10_000_000_000.0)
    assert repartition["produits_de_base"] == 0


PERTES_D_ESSAI = [
    # Deux fraudes internes, pour vérifier le cumul et la perte maximale.
    {"cause_racine": "Fraude interne", "perte_brute": 30_000_000.0, "date_comptabilisation": "2025-03-04"},
    {"cause_racine": "Fraude interne", "perte_brute": 12_000_000.0, "date_comptabilisation": "2025-09-18"},
    {"cause_racine": "Fraude externe", "perte_brute": 8_000_000.0, "date_comptabilisation": "2025-05-22"},
    {"cause_racine": "Défaillance système", "perte_brute": 5_000_000.0, "date_comptabilisation": "2025-11-02"},
    # Cause ambiguë : versée à la catégorie résiduelle, et signalée.
    {"cause_racine": "Erreur humaine", "perte_brute": 3_000_000.0, "date_comptabilisation": "2025-02-01"},
    # Hors exercice : ne doit pas être déclarée.
    {"cause_racine": "Fraude interne", "perte_brute": 99_000_000.0, "date_comptabilisation": "2026-01-15"},
    # Sans date de comptabilisation : la date de survenance prend le relais.
    {"cause_racine": "Fraude externe", "perte_brute": 2_000_000.0, "date_occurrence": "2025-07-07"},
]


@pytest.fixture
def ep22(monkeypatch):
    """EP22 renseigné à partir d'un jeu de pertes maîtrisé."""

    monkeypatch.setattr(
        "app.rapports.fodep.service._lire_pertes_operationnelles",
        lambda: [dict(perte) for perte in PERTES_D_ESSAI],
    )
    classeur = load_workbook(CHEMIN_MODELE)
    anomalies = _remplir_ep22(classeur, date(2026, 6, 30))
    feuille = classeur["EP22"]
    lignes = indexer_codes_dispru(feuille)

    def montants(code: str) -> dict[str, float]:
        ligne = lignes[code]
        colonnes = ("nombre", "total", "maximum", "cinq_plus_grandes")
        return {
            nom: feuille.cell(row=ligne, column=colonne).value or 0
            for nom, colonne in zip(colonnes, range(3, 7))
        }

    return montants, anomalies


def test_l_ep22_ventile_les_pertes_par_categorie_d_evenement(ep22):
    montants, _ = ep22

    fraude_interne = montants("RO011")
    assert fraude_interne["nombre"] == 2, "La perte hors exercice ne doit pas compter."
    assert fraude_interne["total"] == en_millions(42_000_000.0)
    assert fraude_interne["maximum"] == en_millions(30_000_000.0)

    # Deux fraudes externes : l'une datée par sa comptabilisation, l'autre par
    # sa survenance faute de mieux.
    assert montants("RO012")["nombre"] == 2
    assert montants("RO016")["nombre"] == 1, "Défaillance système → interruptions."
    assert montants("RO017")["nombre"] == 1, "Cause ambiguë → catégorie résiduelle."


def test_le_total_de_l_ep22_egale_la_somme_de_ses_categories(ep22):
    """Un état dont le total ne suit pas ses lignes fait rejeter la déclaration."""

    montants, _ = ep22
    total = montants("RO018")
    somme = sum(montants(code)["total"] for code in CODES_EP22)
    nombre = sum(montants(code)["nombre"] for code in CODES_EP22)

    assert total["nombre"] == nombre == 6
    assert abs(total["total"] - somme) <= TOLERANCE_ARRONDI
    assert total["maximum"] == en_millions(30_000_000.0)


def test_l_ep22_signale_les_classifications_deduites(ep22):
    """La convention de classement doit rester visible avant transmission."""

    _, anomalies = ep22
    assert any("Exécution des opérations" in anomalie for anomalie in anomalies), (
        "Une perte rangée par convention plutôt que par déclaration doit être "
        "signalée : sinon la classification passe pour une donnée déclarée."
    )


@pytest.mark.parametrize(
    "arrete, exercice",
    [
        (date(2026, 6, 30), 2025),
        (date(2026, 12, 31), 2026),
        (date(2026, 1, 1), 2025),
    ],
)
def test_l_exercice_declare_est_le_dernier_exercice_clos(arrete, exercice):
    debut, fin = _exercice_declare(arrete)
    assert (debut, fin) == (date(exercice, 1, 1), date(exercice, 12, 31))


PARTICIPATIONS_D_ESSAI = [
    ("Banque Atlantique CI", "etablissement", 50e9, 6e9, 5.5e9),
    ("NSIA Assurances", "assurance", 30e9, 3e9, 2.8e9),
    ("SICAV Sahel", "autre_financiere", 12e9, 1e9, 1.0e9),
    ("SCI Plateau", "societe_immobiliere", 8e9, 900e6, 800e6),
    ("Sucrivoire SA", "entite_commerciale", 40e9, 9e9, 8.5e9),
    ("Cimenterie du Golfe", "entite_commerciale", 25e9, 2e9, 1.9e9),
]


@pytest.fixture(scope="module")
def classeur_avec_participations():
    """Export produit avec un portefeuille de participations maîtrisé."""

    participations = [
        ParticipationView(
            id=index,
            denomination=denomination,
            categorie=categorie,
            capital_entreprise=capital,
            montant_brut=brut,
            montant_net=net,
            commentaire=None,
            cree_le="2026-01-01",
            modifie_le="2026-01-01",
        )
        for index, (denomination, categorie, capital, brut, net) in enumerate(
            PARTICIPATIONS_D_ESSAI, start=1
        )
    ]
    with mock.patch(
        "app.rapports.fodep.service.lister_participations", return_value=participations
    ):
        resultat = construire_fodep()
    return load_workbook(BytesIO(resultat.contenu))


def test_l_ep34_ventile_les_participations_par_section(classeur_avec_participations):
    """Chaque section de l'EP34 doit porter ses propres participations."""

    feuille = classeur_avec_participations["EP34"]
    lignes = indexer_codes_dispru(feuille)

    def denomination(code: str) -> str:
        return feuille.cell(row=lignes[code], column=2).value

    assert denomination("PA001") == "Banque Atlantique CI"
    assert denomination("PA022") == "NSIA Assurances"
    assert denomination("PA043") == "SICAV Sahel"
    assert denomination("PA064") == "SCI Plateau"
    assert denomination("PA085") == "Sucrivoire SA"


def test_le_total_de_l_ep34_egale_la_somme_de_ses_sections(classeur_avec_participations):
    feuille = classeur_avec_participations["EP34"]
    lignes = indexer_codes_dispru(feuille)

    def net(code: str) -> float:
        return feuille.cell(row=lignes[code], column=5).value or 0

    sections = ("PA021", "PA042", "PA063", "PA084", "PA105")
    assert abs(net("PA106") - sum(net(code) for code in sections)) <= TOLERANCE_ARRONDI
    # La section E ne retient que les deux entités commerciales.
    assert net("PA105") == en_millions(8.5e9) + en_millions(1.9e9)


def test_l_ep35_rapporte_chaque_participation_au_capital_de_l_emetteur(
    classeur_avec_participations,
):
    """La limite individuelle de 25 % se lit d = b / a, colonne par colonne."""

    feuille = classeur_avec_participations["EP35"]
    lignes = indexer_codes_dispru(feuille)
    ligne = lignes["PA107"]

    capital = feuille.cell(row=ligne, column=4).value
    brut = feuille.cell(row=ligne, column=5).value
    part = feuille.cell(row=ligne, column=7).value

    assert feuille.cell(row=ligne, column=2).value == "Sucrivoire SA"
    assert part == pytest.approx(brut / capital, abs=1e-3)
    assert part <= 0.25, "Une participation au-delà de 25 % doit sauter aux yeux."


def test_l_ep01_mesure_les_limites_sur_les_participations(classeur_avec_participations):
    """Les normes RA006 à RA008 ne doivent plus être « CONFORME » par défaut."""

    feuille = classeur_avec_participations["EP01"]
    lignes = indexer_codes_dispru(feuille)

    observes = {
        code: feuille.cell(row=lignes[code], column=7).value
        for code in ("RA006", "RA007", "RA008")
    }
    assert all(valeur > 0 for valeur in observes.values()), (
        "Des participations sont déclarées : les trois limites doivent porter "
        f"un niveau observé non nul, or {observes}."
    )
    assert observes["RA006"] == pytest.approx(9.0e9 / 40.0e9, abs=1e-3)


def test_aucune_case_de_saisie_ne_reste_vide_sur_un_etat_declare(classeur):
    """Notice § 3.3 : toute cellule non verrouillée doit être renseignée.

    C'est le garde-fou contre deux oublis symétriques : une donnée de
    l'application qui ne serait pas reportée, et une case du gabarit qui
    resterait telle quelle.

    Le contrôle porte sur les colonnes de montants. Les colonnes
    d'identification — nom de contrepartie, numéro Centrale des risques — en
    sont exclues : un zéro n'y remplace rien, il inventerait une contrepartie
    nommée « 0 ». Leur absence est signalée en anomalie, pas comblée.
    """

    declares = sorted(ETATS_ALIMENTES | ETATS_DECLARES_A_ZERO)
    vides: list[str] = []

    for code in declares:
        if code not in classeur.sheetnames:
            continue
        feuille = classeur[code]
        # La mise en forme gonfle `max_column` bien au-delà du tableau : la
        # largeur utile est celle du bloc d'en-tête, seul à porter des
        # libellés de colonnes.
        largeur = max(
            (
                cellule.column
                for ligne in feuille.iter_rows(min_row=1, max_row=11)
                for cellule in ligne
                if cellule.value not in (None, "")
            ),
            default=0,
        )
        premiere = PREMIERE_COLONNE_NUMERIQUE.get(code, 3)
        if largeur < premiere:
            continue
        for ligne in feuille.iter_rows(min_row=9, min_col=premiere, max_col=largeur):
            for cellule in ligne:
                if isinstance(cellule, MergedCell) or cellule.value is not None:
                    continue
                if cellule.protection is not None and not cellule.protection.locked:
                    vides.append(f"{code}!{cellule.coordinate}")

    assert not vides, (
        "Ces cases de saisie sont restées vides alors que leur état est "
        f"déclaré : {', '.join(vides[:20])}"
        + (f" (et {len(vides) - 20} autres)" if len(vides) > 20 else "")
    )


def test_completer_etat_a_zero_ne_materialise_pas_la_grille_vide():
    """Une borne de feuille gonflée ne doit pas créer ses cellules absentes."""

    classeur = Workbook()
    feuille = classeur.active
    feuille["C9"].protection = Protection(locked=False)
    # Simule l'artefact du modèle : une cellule vide très loin du tableau fixe
    # max_column, sans que la grille intermédiaire n'existe dans le fichier.
    feuille.cell(row=1, column=1025)
    cellules_avant = len(feuille._cells)

    completer_etat_a_zero(feuille)

    assert feuille["C9"].value == 0
    assert len(feuille._cells) == cellules_avant


POSITIONS_MARCHE_D_ESSAI = {
    "actions": {
        "longues": 12.0e9,
        "courtes": 3.0e9,
        "nettes_longues": 10.0e9,
        "nettes_courtes": 1.0e9,
        "base_specifique_liquide": 4.0e9,
        "base_specifique_autre": 7.0e9,
        "base_generale": 9.0e9,
    },
    "change": {
        "longues": 8.0e9,
        "courtes": 5.0e9,
        "position_nette_globale": 8.0e9,
    },
    "taux": {
        "specifique": [
            # Titres d'États de l'UMOA en FCFA : pondération nulle.
            {
                "categorie": "uemoaSovereignXof",
                "qualite": "unrated",
                "tranche": "long",
                "longues": 30.0e9,
                "courtes": 0.0,
            },
            # Souverain A+ à BBB-, durée moyenne : 1 %.
            {
                "categorie": "sovereignDebt",
                "qualite": "aToBbb",
                "tranche": "moyen",
                "longues": 8.0e9,
                "courtes": 1.0e9,
            },
            # Autre émetteur non noté : 8 %.
            {
                "categorie": "otherDebt",
                "qualite": "unrated",
                "tranche": "long",
                "longues": 4.0e9,
                "courtes": 0.0,
            },
        ],
        "general": {
            "tranches": [
                {"index": 0, "zone": 1, "nettes_longues": 2.0e9, "nettes_courtes": 0.5e9},
                {"index": 4, "zone": 2, "nettes_longues": 3.0e9, "nettes_courtes": 1.0e9},
                {"index": 9, "zone": 3, "nettes_longues": 1.0e9, "nettes_courtes": 2.0e9},
            ],
            "equilibre_toutes_tranches": 3.5e9,
            "equilibre_plage_1": 0.5e9,
            "equilibre_plage_2": 1.0e9,
            "equilibre_plage_3": 1.0e9,
            "equilibre_plages_1_2": 0.4e9,
            "equilibre_plages_2_3": 0.3e9,
            "equilibre_plages_1_3": 0.2e9,
            "residu": 0.9e9,
        },
    },
}


@pytest.fixture(scope="module")
def classeur_avec_positions_marche():
    """Export produit avec des positions de marché maîtrisées."""

    risque_marche = {
        "rwa_marche": 194_688_063_683.7,
        "exigence_fonds_propres": 15.6e9,
        "ventilation": {
            "taux": 6.6e9,
            "actions": 1.4e9,
            "change": 0.64e9,
            "produits_de_base": 0.0,
        },
        "positions": POSITIONS_MARCHE_D_ESSAI,
    }
    with mock.patch(
        "app.rapports.fodep.service._lire_risque_marche", return_value=risque_marche
    ):
        resultat = construire_fodep()
    return load_workbook(BytesIO(resultat.contenu))


PARTIES_LIEES_D_ESSAI = [
    {"nom": "Actionnaire de référence", "categorie": "actionnaire"},
    {"nom": "Directeur général", "categorie": "organe_executif"},
    {"nom": "Agent de guichet", "categorie": "personnel_execution"},
]

IMMOBILISATIONS_D_ESSAI = SyntheseImmobilisations(
    exploitation_brut=4.0e9,
    exploitation_net=3.0e9,
    hors_exploitation_brut=1.5e9,
    hors_exploitation_net=1.2e9,
)


@pytest.fixture(scope="module")
def classeur_avec_encours():
    """Export produit avec des immobilisations et des parties liées."""

    def agrege(bilan: float, hors_bilan: float) -> dict:
        """Reproduit la forme complète d'un agrégat par contrepartie.

        Les états EP29 à EP32 consomment les mêmes agrégats : leur fournir une
        forme partielle les ferait échouer sur une clé absente.
        """

        return {
            "pays": "Côte d'Ivoire",
            "secteur": "Administration",
            "bilan": bilan,
            "hors_bilan": hors_bilan,
            "souffrance": 0.0,
            "provisions": 0.0,
            "encours_brut": bilan,
            "apr": 0.0,
            "echeances": defaultdict(float),
            "hors_division": True,
            "contrepartie": False,
        }

    groupes = {
        "Actionnaire de référence": agrege(2.0e9, 0.5e9),
        "Directeur général": agrege(0.4e9, 0.0),
        "Agent de guichet": agrege(0.1e9, 0.0),
    }
    reel = construire_fodep.__globals__["_agreger_par_contrepartie"]

    def agregation_enrichie(expositions):
        resultat = reel(expositions)
        resultat.update(groupes)
        return resultat

    with mock.patch(
        "app.rapports.fodep.service._lire_immobilisations",
        return_value=IMMOBILISATIONS_D_ESSAI,
    ), mock.patch(
        "app.rapports.fodep.service._lire_parties_liees",
        return_value=PARTIES_LIEES_D_ESSAI,
    ), mock.patch(
        "app.rapports.fodep.service._agreger_par_contrepartie",
        side_effect=agregation_enrichie,
    ):
        resultat = construire_fodep()
    return load_workbook(BytesIO(resultat.contenu))


def test_l_ep30_declare_aussi_le_hors_bilan():
    """Une contrepartie sans hors-bilan contredirait les EP29, EP31 et EP32.

    Les quatre états décrivent la même population : omettre les engagements
    par signature ici ferait apparaître un client dont le risque paraîtrait
    plus faible qu'il ne l'est ailleurs dans le formulaire.
    """

    membres = [
        {
            "numero_groupe": "GR-01",
            "nom_groupe": "Groupe Alpha",
            "numero_contrepartie": "CR-100",
            "categorie_lien": "controle_de_droit",
            "nom": "SOCIETE ALPHA",
            "pays": "Côte d'Ivoire",
            "secteur": "Agro-industrie",
        }
    ]
    groupes = {
        "SOCIETE ALPHA": {
            "bilan": 6.0e9,
            "hors_bilan": 2.0e9,
            "souffrance": 0.5e9,
            "provisions": 0.0,
            "encours_brut": 6.0e9,
            "apr": 0.0,
            "echeances": defaultdict(float),
            "pays": "Côte d'Ivoire",
            "secteur": "Agro-industrie",
            "hors_division": False,
            "contrepartie": True,
        }
    }

    classeur = load_workbook(CHEMIN_MODELE)
    with mock.patch(
        "app.rapports.fodep.service._lire_membres_de_groupes", return_value=membres
    ):
        _remplir_ep30(classeur, groupes)

    feuille = classeur["EP30"]
    lignes = indexer_codes_dispru(feuille)
    ligne = lignes[sorted(code for code in lignes if code.startswith("GR"))[0]]
    montant = lambda colonne: feuille.cell(row=ligne, column=colonne).value or 0

    assert montant(10) == en_millions(6.0e9), "Bilan porté en prêts et avances."
    assert montant(13) == en_millions(2.0e9), (
        "Le hors-bilan doit être déclaré en engagements de financement."
    )
    assert montant(9) == en_millions(0.5e9), "Dont en souffrance."


def test_l_ep37_totalise_immobilisations_et_participations(classeur_avec_encours):
    """Le total doit être la somme de ses deux composantes déclarées."""

    feuille = classeur_avec_encours["EP37"]
    lignes = indexer_codes_dispru(feuille)
    net = lambda code: feuille.cell(row=lignes[code], column=4).value or 0

    assert net("IM007") == en_millions(3.0e9), "Immobilisations d'exploitation."
    assert net("IM004") == en_millions(1.2e9), "Immobilisations hors exploitation."
    assert abs(net("IM008") - (net("IM007") + net("IM004"))) <= TOLERANCE_ARRONDI
    assert abs(net("IM009") - (net("IM008") + net("PA106"))) <= TOLERANCE_ARRONDI


def test_l_ep38_ventile_les_concours_par_categorie_de_beneficiaire(
    classeur_avec_encours,
):
    """Les huit colonnes du formulaire doivent totaliser la colonne TOTAL."""

    feuille = classeur_avec_encours["EP38"]
    lignes = indexer_codes_dispru(feuille)

    for code in ("PR001", "PR002", "PR003"):
        ligne = lignes[code]
        categories = sum(
            feuille.cell(row=ligne, column=colonne).value or 0
            for _, colonne in COLONNES_PARTIES_LIEES
        )
        total = feuille.cell(row=ligne, column=11).value or 0
        assert abs(total - categories) <= TOLERANCE_ARRONDI, (
            f"EP38 {code} : total de {total} pour {categories} en colonnes."
        )

    # L'actionnaire porte 2 Md au bilan et 0,5 Md par signature.
    assert (feuille.cell(row=lignes["PR001"], column=3).value or 0) == en_millions(2.0e9)
    assert (feuille.cell(row=lignes["PR002"], column=3).value or 0) == en_millions(0.5e9)


def test_l_ep39_coche_la_colonne_du_beneficiaire(classeur_avec_encours):
    """Un nom sans colonne cochée ne dirait pas à quel titre il figure."""

    feuille = classeur_avec_encours["EP39"]
    lignes = indexer_codes_dispru(feuille)
    codes = sorted(code for code in lignes if code.startswith("PR"))
    colonnes = dict(COLONNES_PARTIES_LIEES)

    for code, partie in zip(codes, PARTIES_LIEES_D_ESSAI):
        ligne = lignes[code]
        assert feuille.cell(row=ligne, column=2).value == partie["nom"]
        colonne = colonnes[partie["categorie"]]
        assert feuille.cell(row=ligne, column=colonne).value == "X", (
            f"EP39 {code} : la colonne « {partie['categorie']} » doit être cochée."
        )


def test_l_ep01_mesure_les_limites_sur_encours(classeur_avec_encours):
    """Les trois dernières normes ne doivent plus être conformes par défaut."""

    feuille = classeur_avec_encours["EP01"]
    lignes = indexer_codes_dispru(feuille)
    observe = lambda code: feuille.cell(row=lignes[code], column=7).value or 0

    assert observe("RA009") > 0, "Limite sur les immobilisations hors exploitation."
    assert observe("RA010") > 0, "Limite sur immobilisations et participations."
    assert observe("RA011") > 0, "Limite sur les prêts aux parties liées."


def test_l_ep25_classe_chaque_titre_sur_la_ligne_de_sa_categorie(
    classeur_avec_positions_marche,
):
    """La catégorie d'émetteur et la notation décident de la ligne.

    Un titre d'État de l'UMOA en FCFA est pondéré à 0 % ; le même montant sur
    un émetteur non noté l'est à 8 %. Les confondre fausserait l'exigence.
    """

    feuille = classeur_avec_positions_marche["EP25"]
    lignes = indexer_codes_dispru(feuille)
    assiette = lambda code: feuille.cell(row=lignes[code], column=7).value or 0

    assert assiette("RM002") == en_millions(30.0e9), "Titres UMOA en FCFA."
    assert assiette("RM004") == en_millions(9.0e9), "Souverain A+/BBB-, moyen terme."
    assert assiette("RM014") == en_millions(4.0e9), "Autre émetteur non noté."


def test_l_ep25_applique_les_facteurs_de_compensation_du_formulaire(
    classeur_avec_positions_marche,
):
    """Chaque compensation vaut la position équilibrée fois son facteur.

    Les facteurs — 10 % par tranche, 40/30/30 % par plage, 40/40/100 % entre
    plages — sont imprimés par le formulaire dans une colonne verrouillée.
    """

    feuille = classeur_avec_positions_marche["EP25"]
    lignes = indexer_codes_dispru(feuille)

    total = 0.0
    for code in ("RM031", "RM032", "RM033", "RM034", "RM035", "RM036", "RM037", "RM038"):
        ligne = lignes[code]
        equilibre = feuille.cell(row=ligne, column=7).value or 0
        facteur = feuille.cell(row=ligne, column=8).value or 0
        exigence = feuille.cell(row=ligne, column=9).value or 0
        assert abs(exigence - equilibre * facteur) <= TOLERANCE_ARRONDI, (
            f"EP25 {code} : exigence de {exigence} pour {equilibre} équilibré "
            f"au facteur {facteur}."
        )
        total += exigence

    declare = feuille.cell(row=lignes["RM039"], column=9).value or 0
    assert abs(declare - total) <= TOLERANCE_ARRONDI, (
        f"Le total du risque général déclare {declare} pour {total} de lignes."
    )


def test_l_ep25_reporte_l_echelle_de_maturite_dans_ses_zones(
    classeur_avec_positions_marche,
):
    """Chaque en-tête de zone doit cumuler les tranches qu'elle couvre."""

    feuille = classeur_avec_positions_marche["EP25"]
    lignes = indexer_codes_dispru(feuille)

    # Zone 1 ne porte ici qu'une tranche : 2 Md longs, 0,5 Md courts.
    assert (feuille.cell(row=lignes["RM125"], column=3).value or 0) == en_millions(2.0e9)
    assert (feuille.cell(row=lignes["RM125"], column=4).value or 0) == en_millions(0.5e9)
    # La tranche « 0 ≤ 1 mois » porte les mêmes montants en positions nettes.
    assert (feuille.cell(row=lignes["RM016"], column=5).value or 0) == en_millions(2.0e9)
    assert (feuille.cell(row=lignes["RM016"], column=6).value or 0) == en_millions(0.5e9)


def test_l_ep26_applique_la_ponderation_imprimee_par_le_formulaire(
    classeur_avec_positions_marche,
):
    """L'exigence doit valoir l'assiette fois la pondération affichée à côté.

    C'est le contrôle le plus immédiat qu'un lecteur puisse faire : la
    pondération est imprimée sur la même ligne, dans une colonne que la BCEAO
    verrouille.
    """

    feuille = classeur_avec_positions_marche["EP26"]
    lignes = indexer_codes_dispru(feuille)

    for code in ("RM047", "RM048", "RM050", "RM051"):
        ligne = lignes[code]
        assiette = feuille.cell(row=ligne, column=7).value or 0
        ponderation = feuille.cell(row=ligne, column=8).value or 0
        exigence = feuille.cell(row=ligne, column=9).value or 0
        assert abs(exigence - assiette * ponderation) <= TOLERANCE_ARRONDI, (
            f"EP26 {code} : exigence de {exigence} pour une assiette de "
            f"{assiette} pondérée à {ponderation}."
        )


def test_les_totaux_de_l_ep26_somment_leurs_lignes(classeur_avec_positions_marche):
    feuille = classeur_avec_positions_marche["EP26"]
    lignes = indexer_codes_dispru(feuille)
    exigence = lambda code: feuille.cell(row=lignes[code], column=9).value or 0

    assert abs(exigence("RM049") - (exigence("RM047") + exigence("RM048"))) <= TOLERANCE_ARRONDI
    assert abs(exigence("RM052") - (exigence("RM050") + exigence("RM051"))) <= TOLERANCE_ARRONDI


def test_l_ep27_porte_l_exigence_sur_le_sens_retenu(classeur_avec_positions_marche):
    """L'exigence de change porte sur le plus élevé des deux sens, pas leur somme.

    Un établissement long de 8 et court de 5 ne risque que sur 8 : déclarer 13
    surestimerait l'exigence de plus de moitié.
    """

    feuille = classeur_avec_positions_marche["EP27"]
    lignes = indexer_codes_dispru(feuille)
    ligne = lignes["RM060"]

    longues = feuille.cell(row=ligne, column=3).value or 0
    courtes = feuille.cell(row=ligne, column=4).value or 0
    soumise_longue = feuille.cell(row=ligne, column=7).value or 0
    soumise_courte = feuille.cell(row=ligne, column=8).value or 0
    ponderation = feuille.cell(row=ligne, column=9).value or 0
    exigence = feuille.cell(row=ligne, column=10).value or 0

    assert soumise_longue + soumise_courte == max(longues, courtes)
    assert abs(exigence - max(longues, courtes) * ponderation) <= TOLERANCE_ARRONDI
    assert abs((feuille.cell(row=lignes["RM062"], column=10).value or 0) - exigence) <= TOLERANCE_ARRONDI


def test_les_ratios_de_solvabilite_rapportent_les_fonds_propres_aux_apr(classeur):
    feuille = classeur["EP02"]
    lignes = indexer_codes_dispru(feuille)

    def valeur(code: str) -> float:
        return feuille.cell(row=lignes[code], column=6).value or 0

    apr = valeur("APR04")
    if apr <= 0:
        pytest.skip("Aucun actif pondéré dans ce jeu de données.")
    for code_ratio, code_fonds_propres in (
        ("RA001", "FPI22 / FPC22"),
        ("RA002", "FPI29 / FPC29"),
        ("RA003", "FPI41 / FPC41"),
    ):
        attendu = valeur(code_fonds_propres) / apr
        assert valeur(code_ratio) == pytest.approx(attendu, abs=1e-3)


def test_l_ep01_reprend_les_niveaux_observes_de_l_ep02(classeur):
    """L'état de conformité doit juger les ratios effectivement déclarés."""

    ep01 = classeur["EP01"]
    ep02 = classeur["EP02"]
    lignes_ep01 = indexer_codes_dispru(ep01)
    lignes_ep02 = indexer_codes_dispru(ep02)
    for code in ("RA001", "RA002", "RA003"):
        observe = ep01.cell(row=lignes_ep01[code], column=7).value
        declare = ep02.cell(row=lignes_ep02[code], column=6).value
        assert observe == declare


def test_l_ep03_totalise_les_fonds_propres_effectifs(classeur):
    feuille = classeur["EP03"]
    lignes = indexer_codes_dispru(feuille)

    def montant(code: str) -> float:
        return feuille.cell(row=lignes[code], column=3).value or 0

    assert montant("FPI29") == pytest.approx(montant("FPI22") + montant("FPI28"), abs=1)
    assert montant("FPI41") == pytest.approx(montant("FPI29") + montant("FPI40"), abs=1)


def test_l_ep20_pondere_chaque_nature_d_actif(classeur):
    feuille = classeur["EP20"]
    lignes = indexer_codes_dispru(feuille)
    total_expositions = total_apr = 0.0
    for code, ligne in lignes.items():
        if code == "RC276":
            continue
        exposition = feuille.cell(row=ligne, column=3).value or 0
        ponderation = float(feuille.cell(row=ligne, column=4).value or 0)
        apr = feuille.cell(row=ligne, column=5).value or 0
        assert abs(apr - exposition * ponderation) <= 1
        total_expositions += exposition
        total_apr += apr
    assert abs((feuille.cell(row=lignes["RC276"], column=3).value or 0) - total_expositions) <= TOLERANCE_ARRONDI
    assert abs((feuille.cell(row=lignes["RC276"], column=5).value or 0) - total_apr) <= TOLERANCE_ARRONDI


def test_le_fodep_declare_les_memes_chiffres_que_l_application(classeur):
    """Le formulaire et le tableau de bord doivent dire la même chose.

    C'est l'invariant qui compte le plus : un utilisateur qui lit 700 Md de
    RWA à l'écran et en déclare 730 à la BCEAO ne peut pas savoir lequel est
    faux. Les écarts historiques venaient de trois causes — provisions non
    déduites de l'EAD, multiplicateur d'APR opérationnel divergent, effet
    d'atténuation non plafonné — toutes trois corrigées.
    """

    from app.dashboard.services import get_dashboard_snapshot

    instantane = get_dashboard_snapshot()
    metriques = {metrique.key: metrique.value for metrique in instantane.metrics}
    ep08 = classeur["EP08"]
    ep02 = classeur["EP02"]
    lignes_ep08 = indexer_codes_dispru(ep08)
    lignes_ep02 = indexer_codes_dispru(ep02)

    def apr(code: str) -> float:
        return ep08.cell(row=lignes_ep08[code], column=5).value or 0

    def fonds_propres(code: str) -> float:
        return ep02.cell(row=lignes_ep02[code], column=6).value or 0

    # Tolérance d'un million : le FODEP arrondit chaque montant au million
    # (§ 2.3), l'application non.
    for code_fodep, cle_metrique in (
        ("APR01", "rwa_credit"),
        ("APR02", "rwa_market"),
        ("APR03", "rwa_op"),
        ("APR04", "rwa"),
    ):
        attendu = metriques[cle_metrique] / 1e6
        assert abs(apr(code_fodep) - attendu) <= 1, (
            f"{code_fodep} déclare {apr(code_fodep):,.0f} M FCFA quand "
            f"l'application affiche {attendu:,.0f} M FCFA."
        )

    detail = instantane.fonds_propres
    for code_fodep, montant in (
        ("FPI22 / FPC22", detail.cet1),
        ("FPI29 / FPC29", detail.tier1),
        ("FPI41 / FPC41", detail.total_fp),
    ):
        assert abs(fonds_propres(code_fodep) - montant / 1e6) <= 1


def test_l_ep33_totalise_ses_quatre_sections(classeur):
    """L'exposition de levier est la somme de ses composantes déclarées."""

    feuille = classeur["EP33"]
    lignes = indexer_codes_dispru(feuille)

    def montant(code: str) -> float:
        return feuille.cell(row=lignes[code], column=3).value or 0

    total = montant("RL004") + montant("RL007") + montant("RL010") + montant("RL013")
    assert abs(montant("RL015") - total) <= TOLERANCE_ARRONDI


def test_l_ep31_ventile_chaque_exposition_par_echeance(classeur):
    """Le total d'une ligne doit être celui de ses catégories d'échéances."""

    feuille = classeur["EP31"]
    lignes = indexer_codes_dispru(feuille)
    premiere_tranche = 4
    nombre_de_tranches = len(BORNES_ECHEANCE_EP31) + 2
    colonne_total = premiere_tranche + nombre_de_tranches

    # La grille du formulaire doit correspondre aux bornes que le code applique.
    entete_total = feuille.cell(row=9, column=colonne_total).value
    assert str(entete_total or "").strip().upper() == "TOTAL", (
        "Les catégories d'échéances de l'EP31 ne correspondent plus aux bornes "
        "codées : la colonne de total n'est pas là où le code l'attend."
    )

    codes = [f"GR{index:03d}" for index in range(173, 194)]
    cumul = 0
    for code in codes[:-1]:
        ligne = lignes[code]
        tranches = sum(
            feuille.cell(row=ligne, column=premiere_tranche + index).value or 0
            for index in range(nombre_de_tranches)
        )
        total = feuille.cell(row=ligne, column=colonne_total).value or 0
        assert abs(total - tranches) <= 1, (
            f"EP31 {code} : total de {total} pour des tranches à {tranches}."
        )
        cumul += total

    total_declare = feuille.cell(row=lignes[codes[-1]], column=colonne_total).value or 0
    assert abs(total_declare - cumul) <= TOLERANCE_ARRONDI


def test_l_ep32_recense_les_cinquante_plus_gros_engagements(classeur):
    feuille = classeur["EP32"]
    lignes = indexer_codes_dispru(feuille)
    colonne_net, colonne_hors_bilan, colonne_total = 8, 9, 10

    codes = [f"GR{index:03d}" for index in range(194, 245)]
    assert len(codes) - 1 == 50

    cumul = 0
    renseignes = 0
    for code in codes[:-1]:
        ligne = lignes[code]
        net = feuille.cell(row=ligne, column=colonne_net).value or 0
        hors_bilan = feuille.cell(row=ligne, column=colonne_hors_bilan).value or 0
        total = feuille.cell(row=ligne, column=colonne_total).value or 0
        assert abs(total - (net + hors_bilan)) <= 1
        cumul += total
        if feuille.cell(row=ligne, column=3).value:
            renseignes += 1

    total_declare = feuille.cell(row=lignes[codes[-1]], column=colonne_total).value or 0
    assert abs(total_declare - cumul) <= TOLERANCE_ARRONDI
    assert renseignes <= 50


def test_les_engagements_sont_classes_par_ordre_decroissant(classeur):
    """Un état des « plus gros » engagements qui ne serait pas trié n'en est pas un."""

    feuille = classeur["EP32"]
    lignes = indexer_codes_dispru(feuille)
    montants = [
        feuille.cell(row=lignes[f"GR{index:03d}"], column=10).value or 0
        for index in range(194, 244)
    ]
    renseignes = [montant for montant in montants if montant]
    assert renseignes == sorted(renseignes, reverse=True)


def test_le_perimetre_est_lu_dans_le_formulaire(classeur):
    """La feuille « Liste des états à renseigner » décide du périmètre.

    Les états de la base consolidée ne concernent pas une déclaration
    individuelle : les coder en dur les figerait, alors que la BCEAO peut
    faire évoluer ce découpage d'une version du formulaire à l'autre.
    """

    individuelle = lire_etats_requis(classeur, base="individuelle")
    consolidee = lire_etats_requis(classeur, base="consolidee")

    assert {"EP01", "EP02", "EP03", "EP08", "EP33"} <= individuelle
    # Fonds propres consolidés et dispositions transitoires associées.
    for etat in ("EP05", "EP06", "EP07"):
        assert etat not in individuelle
        assert etat in consolidee
    assert "EP03" not in consolidee

    # Les postes pour mémoire suivent l'état auquel ils se rattachent.
    assert "EP3M" in individuelle and "EP5M" not in individuelle
    assert "EP5M" in consolidee and "EP3M" not in consolidee


def test_une_base_de_declaration_inconnue_est_refusee(classeur):
    with pytest.raises(ValueError, match="Base de déclaration inconnue"):
        lire_etats_requis(classeur, base="individuel")


def test_les_etats_requis_sont_declares_ou_signales(classeur):
    """Aucun état exigé ne peut être laissé vide en silence."""

    resultat = construire_fodep()
    requis = lire_etats_requis(classeur, base=BASE_DE_DECLARATION)
    reserves = " ".join(resultat.anomalies)

    for etat in sorted(requis):
        if etat in ETATS_ALIMENTES or etat in ETATS_DECLARES_A_ZERO:
            continue
        assert etat in reserves, (
            f"{etat} est exigé sur base {BASE_DE_DECLARATION}, n'est pas "
            "renseigné, et n'est pas signalé à l'utilisateur."
        )


def test_les_normes_non_mesurees_de_l_ep01_sont_signalees():
    """Un « CONFORME » déduit d'une absence de données doit être annoncé.

    L'EP01 calcule lui-même la situation de l'établissement : un niveau
    observé nul y devient « CONFORME ». Pour les six normes que
    l'application ne mesure pas, cette conformité n'est pas constatée.

    L'absence de données est construite ici, pas empruntée à la base : dès que
    des participations, des immobilisations et des parties liées y sont saisies,
    ces six normes se mesurent et la réserve n'a plus lieu d'être. Le test
    porterait alors sur l'état du poste de travail, pas sur le comportement.
    """

    with mock.patch(
        "app.rapports.fodep.service.lister_participations", return_value=[]
    ), mock.patch(
        "app.rapports.fodep.service._lire_immobilisations",
        return_value=SyntheseImmobilisations(),
    ), mock.patch(
        "app.rapports.fodep.service._lire_parties_liees", return_value=[]
    ):
        resultat = construire_fodep()
    reserves = " ".join(resultat.anomalies)
    assert "EP01" in reserves
    assert "CONFORME" in reserves
    assert "prêts aux actionnaires" in reserves.lower() or "actionnaires" in reserves


def test_la_date_d_arrete_est_reportee_sur_l_attestation():
    resultat = construire_fodep(date(2026, 5, 31))
    feuille = load_workbook(BytesIO(resultat.contenu))["ADPE"]
    annee = "".join(str(feuille[reference].value) for reference in ("D7", "E7", "F7", "G7"))
    mois = "".join(str(feuille[reference].value) for reference in ("H7", "I7"))
    jour = "".join(str(feuille[reference].value) for reference in ("K7", "L7"))
    assert (annee, mois, jour) == ("2026", "05", "31")
    assert resultat.date_arrete == date(2026, 5, 31)


def test_aucune_case_de_l_attestation_n_attend_un_nombre():
    """Le type d'une case se lit dans son libellé, et l'écran s'y conforme.

    Rien dans l'attestation n'est un montant : ni le nom de l'établissement,
    ni celui des signataires, ni une adresse électronique. Une case déclarée
    « nombre » ferait poser à l'écran un clavier numérique en face d'un
    libellé qui appelle des lettres, et le déclarant ne pourrait pas la
    remplir du tout.
    """

    from app.rapports.fodep.saisies import TYPES_ADPE, catalogue_adpe

    types = {(case.cellule, case.libelle): case.type_saisie for case in catalogue_adpe()}

    assert "nombre" not in types.values()
    assert set(types.values()) <= set(TYPES_ADPE)

    attendus = {
        "Établissement": "texte",
        "Prénoms et nom": "texte",
        "Fonction": "texte",
        "Téléphone": "telephone",
        "Poste": "telephone",
        "E-mail": "email",
        "Date (AAAA-MM-JJ)": "date",
    }
    for (_, libelle), type_saisie in types.items():
        if libelle in attendus:
            assert type_saisie == attendus[libelle], libelle

    # Le CIB et la lettre clé se saisissent caractère par caractère : leur
    # type dit à l'écran de les présenter en grille, et non en ligne.
    assert types[("N7", "CIB — code identifiant bancaire (5 caractères)")] == "code"
    assert types[("T7", "Lettre clé (1 caractère)")] == "code"


def test_le_nom_de_fichier_porte_la_date_d_arrete():
    assert nom_fichier_fodep(date(2026, 5, 31)) == "FODEP_31052026.xlsx"


def test_la_route_sert_le_classeur_et_ses_reserves():
    """Les réserves voyagent en en-tête : elles doivent y tenir en ASCII.

    Un en-tête HTTP ne transporte pas d'accent : un message brut y ferait
    échouer l'envoi de la réponse entière, fichier compris.
    """

    from fastapi.testclient import TestClient

    from app.main import app

    reponse = TestClient(app).get("/rapports/fodep/export?date_arrete=2026-05-31")

    assert reponse.status_code == 200
    assert reponse.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "FODEP_31052026.xlsx" in reponse.headers["content-disposition"]

    entete = reponse.headers["x-fodep-anomalies"]
    entete.encode("ascii")
    assert isinstance(json.loads(entete), list)

    # Le corps doit être un classeur exploitable, pas un flux tronqué.
    assert len(load_workbook(BytesIO(reponse.content)).sheetnames) == 44


@pytest.mark.parametrize(
    ("montant", "attendu"),
    [
        (0, 0),
        (499_999, 0),
        (500_000, 1),
        (1_500_000, 2),
        (-500_000, -1),
        (-1_400_000, -1),
    ],
)
def test_les_montants_sont_arrondis_au_million_le_plus_proche(montant, attendu):
    assert en_millions(montant) == attendu


@pytest.mark.parametrize(
    ("ponderation", "paliers", "attendu"),
    [
        (0.5, (0.0, 0.2, 0.5, 1.0), {0.5: 1.0}),
        (0.75, (0.0, 0.2, 0.5, 1.0), {0.5: 0.5, 1.0: 0.5}),
        (2.5, (0.0, 0.2, 0.5, 1.0), None),
    ],
)
def test_une_ponderation_absente_est_repartie_sur_les_paliers_voisins(
    ponderation, paliers, attendu
):
    """Répartir plutôt qu'arrondir : l'APR reste exact."""

    resultat = repartir_sur_paliers(ponderation, paliers)
    assert resultat == attendu
    if resultat is not None:
        ponderation_reconstituee = sum(
            palier * quote_part for palier, quote_part in resultat.items()
        )
        assert ponderation_reconstituee == pytest.approx(ponderation)


def test_l_ep20_offre_une_ligne_pour_chacun_de_ses_paliers(classeur):
    """La grille de repli doit exister dans le formulaire, aux bons taux."""

    feuille = classeur["EP20"]
    lignes = indexer_codes_dispru(feuille)
    from app.rapports.fodep.agregation import LIGNES_EP20_PAR_PONDERATION

    for ponderation, code in LIGNES_EP20_PAR_PONDERATION.items():
        assert code in lignes, f"Ligne {code} absente de l'EP20."
        declaree = float(feuille.cell(row=lignes[code], column=4).value or 0)
        assert declaree == pytest.approx(ponderation), (
            f"La ligne {code} de l'EP20 est pondérée à {declaree:.0%}, "
            f"pas à {ponderation:.0%}."
        )
    assert PALIERS_EP20 == tuple(sorted(LIGNES_EP20_PAR_PONDERATION))
