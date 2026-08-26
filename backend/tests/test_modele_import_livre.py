"""Le modele d'import livre doit offrir tout ce que l'import sait lire.

`build_template_workbook` sert un fichier statique quand il existe, sans le
regenerer depuis la specification. Ce fichier avait derive : 26 des 30 colonnes
optionnelles de « Template donnees » manquaient, dont `Type_autre_actif`. Sans
elle, aucune nature d'autre actif ne pouvait etre declaree : l'EP36, l'EP37 et
le poste IM011 de l'EP3M restaient vides, et la norme RA009 de l'EP01 se
declarait « CONFORME » faute d'assiette. Deux colonnes *requises* manquaient
aussi sur « CRM_non_financee ».

Ces tests portent sur le classeur que l'utilisateur telecharge, pas sur le
generateur : c'est le fichier servi qui decide de ce qu'on peut declarer.
`modeles_import/completer_modele_import.py --verifier` fait le meme controle en
ligne de commande.
"""

from __future__ import annotations

from io import BytesIO
import warnings

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pytest

from database.services.excel_import_service import (
    EXPECTED_COLUMNS_BY_SHEET,
    EXPECTED_SHEETS,
    OPTIONAL_COLUMNS_BY_SHEET,
    ExcelImportService,
    OTHER_ASSET_TYPE_OPTIONS,
)


@pytest.fixture(scope="module")
def modele_livre():
    service = ExcelImportService.__new__(ExcelImportService)
    with warnings.catch_warnings():
        # openpyxl previent qu'il ne sait pas relire l'extension x14 des
        # validations : sans effet sur les en-tetes et les listes lues ici.
        warnings.simplefilter("ignore")
        contenu = ExcelImportService.build_template_workbook(service)
        return load_workbook(BytesIO(contenu))


def test_le_modele_livre_porte_toutes_les_colonnes_attendues(modele_livre):
    manques: list[str] = []

    for nom_feuille in EXPECTED_SHEETS:
        assert nom_feuille in modele_livre.sheetnames, nom_feuille
        feuille = modele_livre[nom_feuille]
        presentes = {
            feuille.cell(row=1, column=colonne).value
            for colonne in range(1, feuille.max_column + 1)
        }
        requises = list(EXPECTED_COLUMNS_BY_SHEET.get(nom_feuille, ()))
        optionnelles = list(OPTIONAL_COLUMNS_BY_SHEET.get(nom_feuille, ()))
        for colonne in requises + optionnelles:
            if colonne not in presentes:
                marque = "requise" if colonne in requises else "optionnelle"
                manques.append(f"{nom_feuille} : {colonne} ({marque})")

    assert not manques, (
        "Colonnes absentes du modele livre — l'utilisateur ne peut pas les "
        "renseigner :\n" + "\n".join(manques)
    )


def test_la_liste_des_natures_d_autres_actifs_est_complete(modele_livre):
    """La colonne « Type_autre_actif » doit proposer les 13 natures.

    Une nature absente de la liste deroulante est inaccessible a la saisie, et
    Excel refuse la valeur tapee a la main puisque la validation est stricte.
    """

    feuille = modele_livre["Template données"]
    entetes = {
        feuille.cell(row=1, column=colonne).value: colonne
        for colonne in range(1, feuille.max_column + 1)
    }
    col_idx = entetes["Type_autre_actif"]
    col_lettre = get_column_letter(col_idx)

    # Les options vivent soit dans la formule elle-meme, soit sur la feuille
    # technique qu'elle designe (directement ou via plage nommee / INDIRECT).
    proposees: set[str] = set()
    for validation in feuille.data_validations.dataValidation:
        sqref_str = str(validation.sqref or "")
        formule = str(validation.formula1 or "")
        if f"{col_lettre}2:" in sqref_str or f"{col_lettre}:" in sqref_str or "Type_autre_actif" in formule:
            if formule.startswith('"'):
                proposees.update(
                    valeur.strip() for valeur in formule.strip('"').split(",")
                )
            elif "INDIRECT" in formule:
                for def_name in modele_livre.defined_names.values():
                    target = str(getattr(def_name, "attr_text", "") or getattr(def_name, "value", "") or "")
                    if "!" in target:
                        nom_f, plage_str = target.split("!")
                        nom_clean = nom_f.strip("'")
                        if nom_clean in modele_livre.sheetnames:
                            src_ws = modele_livre[nom_clean]
                            for ligne in src_ws[plage_str.replace("$", "")]:
                                for cel in ligne:
                                    if cel.value and str(cel.value).strip() in OTHER_ASSET_TYPE_OPTIONS:
                                        proposees.add(str(cel.value).strip())
            elif "!" in formule:
                nom_feuille, plage = formule.split("!")
                source = modele_livre[nom_feuille.strip("'")]
                for ligne in source[plage.replace("$", "")]:
                    for cellule in ligne:
                        if cellule.value:
                            proposees.add(str(cellule.value).strip())

    absentes = [
        nature for nature in OTHER_ASSET_TYPE_OPTIONS if nature not in proposees
    ]
    assert not absentes, "Natures absentes de la liste deroulante : " + str(absentes)


def test_les_feuilles_de_reference_bceao_sont_conservees(modele_livre):
    """Le modèle doit fournir ses feuilles explicatives et documenter les natures d'actifs."""
    for feuille in ("Instructions", "Listes de référence"):
        assert feuille in modele_livre.sheetnames, feuille

    reference = modele_livre["Listes de référence"]
    libelles = [
        str(reference.cell(row=rang, column=1).value or "")
        + " | "
        + str(reference.cell(row=rang, column=2).value or "")
        for rang in range(1, reference.max_row + 1)
    ]
    joints = " | ".join(libelles)
    assert "Immobilisations hors exploitation" in joints
    assert "Immobilisations incorporelles" in joints
