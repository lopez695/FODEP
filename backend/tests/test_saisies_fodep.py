"""Ce que le declarant saisit se retrouve dans la declaration.

L'attestation ADPE est la seule partie du FODEP que l'application ne calcule
pas : son contenu vient entierement de l'ecran de saisie. Une case perdue en
route ne se verrait qu'a la lecture de la declaration transmise, trop tard.
Ces cas suivent donc chaque champ de bout en bout : ecran, base, classeur.
"""

from __future__ import annotations

from io import BytesIO
import unicodedata

from openpyxl import load_workbook
import pytest

from app.rapports.fodep import construire_fodep
from app.rapports.fodep.saisies import (
    CHAMPS_ADPE,
    CHAMPS_ECLATES,
    PAYS_UEMOA,
    catalogue_adpe,
    enregistrer_saisies,
    lire_saisies,
)
from app.rapports.services import lire_cases_a_saisir


def _valeur_pour(cellule: str, libelle: str, type_champ: str) -> str:
    """Une valeur plausible, propre a la case, pour la reconnaitre a l'arrivee."""

    if type_champ == "pays":
        return "Senegal" if "Senegal" in PAYS_UEMOA else PAYS_UEMOA[6]
    if type_champ == "email":
        return "controle." + cellule.lower() + "@banque.test"
    if type_champ == "telephone":
        return "0022133" + cellule.strip("ABCDEFGHIJKLMNOPQRSTUVWXYZ").zfill(6)
    if type_champ == "date":
        return "2026-03-31"
    if type_champ == "code":
        return "01234" if cellule == "N7" else "K"
    return "Valeur " + cellule


@pytest.fixture
def saisies_restaurees():
    """Rend la table des saisies telle qu'elle etait avant le cas."""

    avant = lire_saisies()
    yield
    apres = lire_saisies()
    enregistrer_saisies(
        [
            {"etat": etat, "cellule": cellule, "valeur": None, "texte": None}
            for (etat, cellule) in apres
        ]
    )
    enregistrer_saisies(
        [
            {
                "etat": etat,
                "cellule": cellule,
                "valeur": saisie.valeur,
                "texte": saisie.texte,
                "commentaire": saisie.commentaire,
            }
            for (etat, cellule), saisie in avant.items()
        ]
    )


def test_chaque_champ_de_l_attestation_survit_a_l_enregistrement(saisies_restaurees):
    attendus = {
        cellule: _valeur_pour(cellule, libelle, type_champ)
        for cellule, libelle, _, type_champ in CHAMPS_ADPE
    }

    touchees = enregistrer_saisies(
        [
            {"etat": "ADPE", "cellule": cellule, "texte": texte}
            for cellule, texte in attendus.items()
        ]
    )
    assert touchees == len(CHAMPS_ADPE)

    # 1. La base rend exactement ce qui lui a ete confie.
    enregistrees = lire_saisies()
    for cellule, texte in attendus.items():
        assert ("ADPE", cellule) in enregistrees, cellule
        assert enregistrees[("ADPE", cellule)].texte == texte, cellule

    # 2. L'ecran les represente : c'est ce que le declarant relit en revenant.
    catalogue = {
        case.cellule: case
        for etat in lire_cases_a_saisir().etats
        for case in etat.cases
        if etat.etat == "ADPE"
    }
    for cellule, texte in attendus.items():
        assert catalogue[cellule].texte == texte, cellule

    # 3. Le classeur transmis les porte, a la cellule que le formulaire prevoit.
    feuille = load_workbook(BytesIO(construire_fodep().contenu))["ADPE"]
    for cellule, texte in attendus.items():
        eclatee = CHAMPS_ECLATES.get(("ADPE", cellule))
        if eclatee:
            ecrit = "".join(str(feuille[case].value or "") for case in eclatee)
        else:
            ecrit = str(feuille[cellule].value or "")
        assert ecrit == texte, cellule


def test_une_case_videe_est_effacee_et_non_conservee(saisies_restaurees):
    enregistrer_saisies([{"etat": "ADPE", "cellule": "T5", "texte": "Banque temoin"}])
    assert lire_saisies()[("ADPE", "T5")].texte == "Banque temoin"

    enregistrer_saisies([{"etat": "ADPE", "cellule": "T5", "texte": "   "}])
    assert ("ADPE", "T5") not in lire_saisies()


def test_l_etat_se_choisit_parmi_les_huit_membres_de_l_union():
    """La case C5 identifie la declaration : elle ne se tape pas.

    Le FODEP est la declaration prudentielle de l'UMOA. Laisser taper le nom
    de l'Etat exposait la case a la faute de frappe et a l'orthographe
    approximative, sur une valeur qui n'a que huit reponses possibles.
    """

    etat = next(case for case in catalogue_adpe() if case.cellule == "C5")

    assert etat.type_saisie == "pays"
    assert tuple(etat.choix) == PAYS_UEMOA
    assert len(PAYS_UEMOA) == 8
    assert "Burkina Faso" in PAYS_UEMOA
    # Ordre alphabetique francais, ou l'accent ne deplace pas la lettre :
    # « Benin » precede « Burkina Faso », ce que le tri par point de code
    # Unicode ferait l'inverse.
    def sans_accent(nom: str) -> str:
        decompose = unicodedata.normalize("NFD", nom)
        return "".join(c for c in decompose if not unicodedata.combining(c)).casefold()

    assert list(PAYS_UEMOA) == sorted(PAYS_UEMOA, key=sans_accent)

    # Les autres cases restent des lignes de saisie libre.
    assert [case.cellule for case in catalogue_adpe() if case.choix] == ["C5"]


def test_la_date_d_arrete_choisie_est_conservee(saisies_restaurees):
    """La date d'arrete se choisit sur l'ecran de saisie : elle s'y enregistre.

    Elle n'etait retenue que le temps de la session. Choisie, puis la page
    rechargee ou l'application rouverte, elle repartait de la date de fin du
    reporting — et l'attestation portait une date que le declarant n'avait pas
    voulue.
    """

    from datetime import date

    from app.rapports.fodep.saisies import lire_date_arrete
    from app.rapports.services import enregistrer_saisies_fodep, lire_cases_a_saisir

    depart = lire_date_arrete()
    try:
        enregistrer_saisies_fodep([], date(2026, 3, 31))
        assert lire_date_arrete() == date(2026, 3, 31)
        assert lire_cases_a_saisir().date_arrete == date(2026, 3, 31)

        # Enregistrer des cases sans toucher a la date ne l'efface pas : sans
        # cela, la premiere saisie suivante la ferait disparaitre.
        enregistrer_saisies_fodep([], None)
        assert lire_date_arrete() == date(2026, 3, 31)

        enregistrer_saisies_fodep([], date(2026, 6, 30))
        assert lire_date_arrete() == date(2026, 6, 30)
    finally:
        if depart is not None:
            enregistrer_saisies_fodep([], depart)


def test_seul_l_adpe_est_expose_a_la_saisie_manuelle():
    """L'écran de saisie manuelle n'expose que l'ADPE (20 champs), rien d'autre."""
    from app.rapports.services import lire_cases_a_saisir

    saisies = lire_cases_a_saisir()
    assert len(saisies.etats) == 1
    assert saisies.etats[0].etat == "ADPE"
    assert len(saisies.etats[0].cases) == 20
    assert saisies.total_cases == 20


# ─── Etats prudentiels sans source ───────────────────────────────────────────


def test_le_catalogue_des_etats_sans_source_suit_le_formulaire():
    """Les cases proposees sont celles que la BCEAO a ouvertes, et elles seules.

    Le catalogue n'enumere rien : il relit le verrouillage du modele, comme le
    reste de l'export. Un ecart entre ce que l'ecran propose et ce que le
    formulaire attend rendrait la saisie inutile -- une case saisie hors d'une
    ligne codee n'a pas d'adresse dans la nomenclature, et la plate-forme ne
    saurait pas la lire.
    """

    from openpyxl.cell.cell import MergedCell

    from app.rapports.fodep.disposition import (
        CHEMIN_MODELE,
        MOTIF_CODE_DISPRU,
        PREMIERE_LIGNE_UTILE,
        styles_ouverts,
    )
    from app.rapports.fodep.saisies import (
        ETATS_FACULTATIFS_CATALOGUE,
        _est_calculee,
        catalogue_etat,
    )

    classeur = load_workbook(CHEMIN_MODELE)
    try:
        styles = styles_ouverts(classeur)
        for nom, _, _note in ETATS_FACULTATIFS_CATALOGUE:
            feuille = classeur[nom]
            attendues = set()
            for cellule in tuple(feuille._cells.values()):
                if (
                    cellule.row < PREMIERE_LIGNE_UTILE
                    or isinstance(cellule, MergedCell)
                    or cellule._style is None
                    or cellule._style.protectionId not in styles
                ):
                    continue
                code = str(feuille.cell(row=cellule.row, column=1).value or "").strip()
                if not MOTIF_CODE_DISPRU.fullmatch(code):
                    continue
                # Les cases que l'export calcule ne se saisissent pas.
                if _est_calculee(nom, code, cellule.column_letter):
                    continue
                attendues.add(cellule.coordinate)
            cases = catalogue_etat(nom)
            assert {case.cellule for case in cases} == attendues, nom
            assert cases, f"{nom} ne propose aucune case"

            # Chaque case se designe comme le formulaire la designe : un code
            # DISPRU et un poste. Sans eux, le declarant ne sait pas ce qu'il
            # remplit.
            for case in cases:
                assert case.code, f"{nom}!{case.cellule} sans code DISPRU"
                assert case.libelle, f"{nom}!{case.cellule} sans libellé"
    finally:
        classeur.close()
def test_une_case_saisie_sur_un_etat_sans_source_arrive_dans_la_declaration():
    """Le zero automatique ne recouvre pas ce que le declarant a porte.

    C'est tout l'objet de ces trois etats : sans saisie ils partent a zero, et
    le formulaire affirme que l'etablissement ne detient ni derive, ni produit
    de base. Avec saisie, la valeur doit survivre au balayage final.
    """

    from app.rapports.fodep.saisies import enregistrer_saisies

    enregistrer_saisies(
        [
            {"etat": "EP28", "cellule": "C12", "valeur": 4200.0},
            {"etat": "EP28", "cellule": "D12", "valeur": 1500.0},
        ]
    )
    try:
        resultat = construire_fodep()
        feuille = load_workbook(BytesIO(resultat.contenu))["EP28"]
        assert feuille["C12"].value == 4200
        assert feuille["D12"].value == 1500
        # La case voisine, laissee vide, retombe bien a zero.
        assert feuille["E12"].value == 0

        # Et l'etat sort de la liste de ceux qu'on declare a zero faute de
        # source : il porte desormais une donnee.
        annonces = [
            reserve.message
            for reserve in resultat.anomalies
            if "déclarés à zéro faute de source" in reserve.message
        ]
        assert annonces and "EP28" not in annonces[0]
    finally:
        enregistrer_saisies(
            [
                {"etat": "EP28", "cellule": "C12", "valeur": None},
                {"etat": "EP28", "cellule": "D12", "valeur": None},
            ]
        )
