"""Analyse d'une declaration FODEP au regard du dispositif prudentiel UMOA.

L'utilisateur depose un PDF de declaration — le sien, celui d'un exercice
precedent, celui d'une filiale — et l'outil dit ou en sont les onze normes de
l'EP01 : respectee, depassee, ou non mesuree.

Le PDF ne sert qu'a lire les *nombres*. Les intitules des normes sont pris dans
le formulaire de la BCEAO lui-meme : le texte extrait d'un PDF arrive colle, sans
espaces (« Limiteindividuellesurlesparticipations... »), et reconstituer les mots
serait une invention. Le formulaire porte deja le libelle exact.

Ce que l'analyse ne sait pas faire : un PDF scanne. Sans couche de texte, il n'y
a rien a extraire, et l'outil n'embarque pas de reconnaissance optique. Le
message le dit plutot que de rendre une declaration vide de normes.
"""

from __future__ import annotations

from functools import lru_cache
from io import BytesIO
import re

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.rapports.fodep.disposition import indexer_codes_dispru
from app.rapports.fodep.reserves import Reserve
from app.rapports.fodep.service import CHEMIN_MODELE, COLONNE_B

# Sens de chaque norme de l'EP01, en dernier recours.
#
# Le sens est normalement lu dans le formulaire lui-meme (voir
# `sens_des_normes`) : c'est la regle de la BCEAO, pas notre lecture du
# dispositif, et elle survivra a une revision du formulaire. Ces ensembles ne
# servent que si la formule devient illisible.
#
#   - minimums : ratios de fonds propres et levier ;
#   - maximums : division des risques et toutes les limites.
MINIMUMS: frozenset[str] = frozenset({"RA001", "RA002", "RA003", "RA005"})
MAXIMUMS: frozenset[str] = frozenset(
    {"RA004", "RA006", "RA007", "RA008", "RA009", "RA010", "RA011"}
)

# Largeur balayee pour l'inventaire : au-dela, les etats du FODEP ne portent
# rien. Meme borne que l'extraction du contenu.
COLONNES_MAX = 14

# Colonnes de l'EP01, telles que le formulaire les dispose.
COLONNE_REFERENCE = 5
COLONNE_SEUIL = 6
COLONNE_OBSERVE = 7
COLONNE_SITUATION = 8

# La formule de la colonne « Situation » : =IF($G10>($F10),"X","Y"). Quand la
# branche vraie porte CONFORME, depasser le seuil est conforme — c'est un
# plancher. Quand elle porte INFRACTION, c'est un plafond.
_MOTIF_SITUATION = re.compile(
    r'IF\s*\(\s*\$?G\d+\s*>\s*\(?\s*\$?F\d+\s*\)?\s*,\s*"([^"]+)"', re.I
)

RESPECTEE = "respectee"
DEPASSEE = "depassee"
NON_MESUREE = "non_mesuree"

# Une norme de l'EP01 dans le texte extrait : son code, son intitule colle, la
# reference de l'etat qui la produit, puis ses nombres.
_MOTIF_NORME = re.compile(r"(RA\d{3})(.*?)(?=RA\d{3}|$)", re.S)
_MOTIF_REFERENCE = re.compile(r"(EP\d[\dA-Za-z]?)")
_MOTIF_NOMBRE = re.compile(r"\d+(?:,\d+)?")


class NormeAnalysee(BaseModel):
    """Une norme prudentielle, et ou elle en est."""

    code: str
    libelle: str

    #: Etat du FODEP qui produit le niveau observe (EP02, EP29, EP35...).
    reference: str = ""

    #: Niveau a respecter, tel que le formulaire le porte.
    seuil: float | None = None

    #: Niveau observe. Absent quand la case est vide — le formulaire la calcule
    #: lui-meme et sa valeur n'existe qu'a l'ouverture du classeur.
    observe: float | None = None

    #: `True` pour un plancher, `False` pour un plafond, `None` si le sens de la
    #: norme est inconnu de l'outil.
    minimum: bool | None = None

    situation: str = NON_MESUREE

    #: Marge restante, ou montant du depassement. En points du ratio.
    ecart: float | None = None


EXACT = "exact"
ARRONDI = "arrondi"
ECART = "ecart"


class ControleCoherence(BaseModel):
    """Un total du formulaire, confronte a la somme de ses lignes.

    Les normes de l'EP01 disent si la declaration respecte le dispositif ; ces
    controles disent si elle se contredit. Un total qui ne suit pas ses lignes
    fait rejeter le depot sans qu'aucune norme ne soit en cause.
    """

    etat: str
    libelle: str

    #: Colonne verifiee, telle que le formulaire l'intitule.
    colonne: str

    #: Somme des lignes de la section.
    attendu: float

    #: Valeur portee par la ligne de total.
    constate: float

    ecart: float

    #: Derive maximale imputable aux arrondis : une demi-unite par ligne sommee.
    tolerance: float

    statut: str = EXACT


class EtatRenseigne(BaseModel):
    """Ce qu'un etat porte : combien de lignes, et s'il est entierement a zero."""

    nom: str
    lignes: int = 0

    #: Aucune valeur numerique non nulle. L'etat est declare, mais ne dit rien —
    #: et un formulaire calcule « CONFORME » sur des zeros comme sur des
    #: encours mesures.
    tout_a_zero: bool = False


class AnalyseDeclaration(BaseModel):
    """Verdict d'ensemble sur la declaration deposee."""

    nom_fichier: str
    pages: int
    normes: list[NormeAnalysee] = Field(default_factory=list)

    #: Controles de coherence interne. Vides pour une impression : sommer des
    #: nombres extraits d'un PDF ne prouverait rien.
    controles: list[ControleCoherence] = Field(default_factory=list)

    #: Etats du formulaire, et ce qu'ils portent.
    inventaire: list[EtatRenseigne] = Field(default_factory=list)

    #: `True` quand la lecture a porte sur le classeur, la piece transmise.
    classeur: bool = False

    #: Ce qui empeche de conclure, s'il y a lieu.
    avertissements: list[str] = Field(default_factory=list)

    #: Les reserves que l'export a formulees en renseignant le formulaire.
    #: Vides pour une declaration deposee : elles naissent du remplissage, pas
    #: de la lecture. C'est ce qui distingue l'analyse de la declaration EN
    #: COURS de celle d'un fichier recu.
    reserves: list[Reserve] = Field(default_factory=list)

    @property
    def depassees(self) -> list[NormeAnalysee]:
        return [n for n in self.normes if n.situation == DEPASSEE]

    @property
    def non_mesurees(self) -> list[NormeAnalysee]:
        return [n for n in self.normes if n.situation == NON_MESUREE]


@lru_cache(maxsize=1)
def _reference_du_formulaire() -> tuple[dict[str, str], dict[str, bool]]:
    """Intitules et sens des normes, lus dans le formulaire de la BCEAO.

    Deux choses que le document analyse ne donne pas :

    - l'intitule exact, parce que le texte d'un PDF arrive colle, sans espaces ;
    - le sens de la norme, parce que le formulaire n'affiche que le niveau a
      respecter. Il est dans la formule de la colonne « Situation ».
    """

    classeur = load_workbook(CHEMIN_MODELE, data_only=False)
    try:
        feuille = classeur["EP01"]
        libelles: dict[str, str] = {}
        sens: dict[str, bool] = {}
        for code, rang in indexer_codes_dispru(feuille).items():
            if not code.startswith("RA"):
                continue
            libelles[code] = str(
                feuille.cell(row=rang, column=COLONNE_B).value or ""
            ).strip()
            formule = str(feuille.cell(row=rang, column=COLONNE_SITUATION).value or "")
            if (trouve := _MOTIF_SITUATION.search(formule)) is not None:
                sens[code] = trouve.group(1).strip().upper().startswith("CONFORME")
        return libelles, sens
    finally:
        classeur.close()


def libelles_des_normes() -> dict[str, str]:
    return _reference_du_formulaire()[0]


def sens_des_normes() -> dict[str, bool]:
    """Code de norme -> `True` pour un plancher, `False` pour un plafond."""

    return _reference_du_formulaire()[1]


def _nombre(texte: str) -> float | None:
    try:
        return float(texte.replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _situation(
    code: str, seuil: float | None, observe: float | None
) -> tuple[str, bool | None, float | None]:
    sens = sens_des_normes()
    minimum: bool | None = sens.get(code)
    if minimum is None:
        # Le formulaire n'a pas livre son sens : on retombe sur le dispositif.
        if code in MINIMUMS:
            minimum = True
        elif code in MAXIMUMS:
            minimum = False

    if observe is None or seuil is None or minimum is None:
        return NON_MESUREE, minimum, None

    if minimum:
        # Plancher : la marge est ce qui depasse le seuil.
        #
        # L'inegalite est stricte, comme dans le formulaire — sa formule declare
        # INFRACTION quand le niveau observe egale exactement le niveau a
        # respecter. Adoucir en « superieur ou egal » serait plus intuitif, mais
        # l'analyse dirait alors le contraire de ce que la declaration declarera.
        ecart = round(observe - seuil, 6)
        return (RESPECTEE if observe > seuil else DEPASSEE), minimum, ecart

    # Plafond : la marge est ce qui reste avant de l'atteindre. Le formulaire y
    # accepte l'egalite — seul un depassement franc est une infraction.
    ecart = round(seuil - observe, 6)
    return (RESPECTEE if observe <= seuil else DEPASSEE), minimum, ecart


def analyser_texte(texte: str) -> list[NormeAnalysee]:
    """Normes lues dans le texte de l'EP01."""

    libelles = libelles_des_normes()
    compact = texte.replace("\n", " ")
    normes: list[NormeAnalysee] = []
    vus: set[str] = set()

    for trouve in _MOTIF_NORME.finditer(compact):
        code = trouve.group(1)
        if code in vus:
            continue
        segment = trouve.group(2)

        reference = ""
        queue = segment
        if (ref := _MOTIF_REFERENCE.search(segment)) is not None:
            reference = ref.group(1)
            # Les nombres se lisent apres la reference : l'intitule en contient
            # lui-meme (« 25 % du capital »), et les confondre ferait passer un
            # rappel de plafond pour un niveau observe.
            queue = segment[ref.end() :]

        nombres = [_nombre(n) for n in _MOTIF_NOMBRE.findall(queue)]
        nombres = [n for n in nombres if n is not None]
        seuil = nombres[0] if nombres else None
        observe = nombres[1] if len(nombres) > 1 else None

        situation, minimum, ecart = _situation(code, seuil, observe)
        normes.append(
            NormeAnalysee(
                code=code,
                libelle=libelles.get(code, ""),
                reference=reference,
                seuil=seuil,
                observe=observe,
                minimum=minimum,
                situation=situation,
                ecart=ecart,
            )
        )
        vus.add(code)

    return sorted(normes, key=lambda norme: norme.code)


def _nombre_ou_none(valeur) -> float | None:
    if isinstance(valeur, bool) or valeur is None:
        return None
    if isinstance(valeur, (int, float)):
        return float(valeur)
    return _nombre(str(valeur))


def analyser_classeur(
    octets: bytes,
) -> tuple[list[NormeAnalysee], list[ControleCoherence], list[EtatRenseigne]]:
    """Normes lues dans un classeur FODEP.

    C'est la lecture exacte : les niveaux y sont des nombres, aux cases que le
    formulaire leur reserve. Aucune extraction de texte, donc aucune ambiguite —
    et c'est le classeur qui est transmis a la BCEAO, pas son impression.
    """

    try:
        classeur = load_workbook(BytesIO(octets), data_only=True)
    except Exception as exc:  # openpyxl leve des types varies sur un fichier casse
        raise ValueError(f"Classeur illisible : {exc}") from exc

    try:
        nom = next(
            (feuille for feuille in classeur.sheetnames if feuille.strip() == "EP01"),
            None,
        )
        if nom is None:
            raise ValueError(
                "Aucun état EP01 dans ce classeur : ce n'est pas une déclaration "
                "FODEP."
            )

        feuille = classeur[nom]
        libelles = libelles_des_normes()
        normes: list[NormeAnalysee] = []
        for code, rang in sorted(indexer_codes_dispru(feuille).items()):
            if not code.startswith("RA"):
                continue
            seuil = _nombre_ou_none(feuille.cell(row=rang, column=COLONNE_SEUIL).value)
            observe = _nombre_ou_none(
                feuille.cell(row=rang, column=COLONNE_OBSERVE).value
            )
            reference = str(
                feuille.cell(row=rang, column=COLONNE_REFERENCE).value or ""
            ).strip()
            situation, minimum, ecart = _situation(code, seuil, observe)
            normes.append(
                NormeAnalysee(
                    code=code,
                    libelle=libelles.get(code, ""),
                    reference=reference,
                    seuil=seuil,
                    observe=observe,
                    minimum=minimum,
                    situation=situation,
                    ecart=ecart,
                )
            )
        return (
            normes,
            _controles_du_classeur(classeur),
            _inventaire_du_classeur(classeur),
        )
    finally:
        classeur.close()


def _lecteur_de_pdf():
    """Charge pypdf a l'usage, pas a l'import du module.

    Une dependance manquante ne doit pas empecher l'API de demarrer : importer
    pypdf en tete de fichier faisait echouer `app.main` en entier, donc *toutes*
    les routes, pour une bibliotheque qui ne sert qu'a lire un PDF. L'analyse du
    classeur, elle, n'en a pas besoin.
    """

    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:
        raise ValueError(
            "Lecture des PDF indisponible : la bibliothèque pypdf n'est pas "
            "installée dans cet environnement. Déposez le classeur (.xlsx) de la "
            "déclaration, ou installez les dépendances du backend "
            "(pip install -r requirements.txt)."
        ) from exc
    return PdfReader, PdfReadError


# Totaux dont la composition ne souffre aucune interpretation : la ligne de
# total additionne exactement les lignes citees, sans deduction.
#
# Les autres totaux du formulaire ne s'y pretent pas. Dans l'EP03, « TOTAL DES
# FONDS PROPRES CET1 » retranche des deductions au lieu d'additionner : une regle
# generale « le total egale la somme des lignes au-dessus » y produirait de faux
# ecarts, et un controle qui crie a tort ne se lit plus.
#
#   etat, code du total, codes sommes, colonne, intitule de la colonne
CONTROLES: tuple[tuple[str, str, tuple[str, ...], int, str], ...] = (
    (
        "EP20",
        "RC276",
        (
            "RC262", "RC263", "RC264", "PA166", "RC265", "RC266", "RC267",
            "RC268", "RC269", "RC270", "RC271", "RC272", "RC273", "PA176",
            "ID012", "RC274", "RC275",
        ),
        3,
        "Exposition nette",
    ),
    (
        "EP20",
        "RC276",
        (
            "RC262", "RC263", "RC264", "PA166", "RC265", "RC266", "RC267",
            "RC268", "RC269", "RC270", "RC271", "RC272", "RC273", "PA176",
            "ID012", "RC274", "RC275",
        ),
        5,
        "Actifs ponderes",
    ),
    (
        "EP34",
        "PA106",
        ("PA021", "PA042", "PA063", "PA084", "PA105"),
        4,
        "Souscription brute",
    ),
    (
        "EP34",
        "PA106",
        ("PA021", "PA042", "PA063", "PA084", "PA105"),
        5,
        "Libere net",
    ),
) + tuple(
    # Chaque section de l'EP34 totalise ses vingt lignes de saisie.
    (
        "EP34",
        code_total,
        tuple(f"PA{numero:03d}" for numero in range(premier, dernier + 1)),
        colonne,
        intitule,
    )
    for code_total, premier, dernier in (
        ("PA021", 1, 20),
        ("PA042", 22, 41),
        ("PA063", 43, 62),
        ("PA084", 64, 83),
        ("PA105", 85, 104),
    )
    for colonne, intitule in ((4, "Souscription brute"), (5, "Libere net"))
)


def _controles_du_classeur(classeur) -> list[ControleCoherence]:
    controles: list[ControleCoherence] = []
    for etat, code_total, codes, colonne, intitule in CONTROLES:
        if etat not in classeur.sheetnames:
            continue
        feuille = classeur[etat]
        lignes = indexer_codes_dispru(feuille)
        if code_total not in lignes:
            continue

        presents = [code for code in codes if code in lignes]
        if not presents:
            continue

        attendu = 0.0
        for code in presents:
            valeur = _nombre_ou_none(
                feuille.cell(row=lignes[code], column=colonne).value
            )
            attendu += valeur or 0.0
        constate = (
            _nombre_ou_none(feuille.cell(row=lignes[code_total], column=colonne).value)
            or 0.0
        )
        ecart = round(constate - attendu, 4)

        # Chaque ligne du formulaire est arrondie a l'unite : la somme peut donc
        # deriver d'une demi-unite par ligne. Au-dela, l'ecart n'est plus
        # imputable a l'arrondi.
        tolerance = max(1.0, len(presents) / 2)
        if ecart == 0:
            statut = EXACT
        elif abs(ecart) <= tolerance:
            statut = ARRONDI
        else:
            statut = ECART

        libelle = str(
            feuille.cell(row=lignes[code_total], column=COLONNE_B).value or code_total
        ).strip()
        controles.append(
            ControleCoherence(
                etat=etat,
                libelle=libelle,
                colonne=intitule,
                attendu=round(attendu, 4),
                constate=round(constate, 4),
                ecart=ecart,
                tolerance=tolerance,
                statut=statut,
            )
        )
    return controles


def _inventaire_du_classeur(classeur) -> list[EtatRenseigne]:
    """Ce que chaque etat porte : combien de lignes, et s'il est tout a zero.

    Un etat entierement a zero est declare, mais ne dit rien. Le lecteur doit le
    voir : le formulaire calcule « CONFORME » sur des zeros comme sur des encours
    mesures.
    """

    inventaire: list[EtatRenseigne] = []
    for nom in classeur.sheetnames:
        feuille = classeur[nom]
        lignes = 0
        des_nombres = False
        valeur_non_nulle = False
        for ligne in feuille.iter_rows(max_col=COLONNES_MAX, values_only=True):
            porteuse = False
            for valeur in ligne:
                if valeur is None or valeur == "":
                    continue
                porteuse = True
                if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
                    des_nombres = True
                    if valeur != 0:
                        valeur_non_nulle = True
            if porteuse:
                lignes += 1
        if lignes:
            inventaire.append(
                EtatRenseigne(
                    nom=nom,
                    lignes=lignes,
                    # Une feuille sans aucun nombre n'est pas « a zero » : elle
                    # est textuelle. La page de garde et l'attestation en sont,
                    # et les signaler ferait passer une mise en page pour une
                    # declaration vide.
                    tout_a_zero=des_nombres and not valeur_non_nulle,
                )
            )
    return inventaire


def _analyser_pdf(octets: bytes) -> tuple[list[NormeAnalysee], list[str], int]:
    PdfReader, PdfReadError = _lecteur_de_pdf()
    try:
        lecteur = PdfReader(BytesIO(octets))
        pages = [page.extract_text() or "" for page in lecteur.pages]
    except (PdfReadError, ValueError, OSError) as exc:
        raise ValueError(f"PDF illisible : {exc}") from exc

    avertissements: list[str] = []
    texte = "\n".join(pages)
    if not texte.strip():
        avertissements.append(
            "Ce PDF ne contient aucun texte : il s'agit probablement d'un scan. "
            "L'outil ne sait pas lire une image, et n'embarque pas de "
            "reconnaissance optique."
        )
    return analyser_texte(texte), avertissements, len(pages)


def analyser_declaration(octets: bytes, nom_fichier: str) -> AnalyseDeclaration:
    """Confronte une declaration deposee au dispositif prudentiel.

    Deux formats, parce que la declaration en a deux : le classeur, qui est la
    piece transmise a la BCEAO, et le PDF, qui en est l'impression. Le format se
    reconnait aux premiers octets plutot qu'a l'extension — un fichier renomme ne
    doit pas faire echouer la lecture.
    """

    controles: list[ControleCoherence] = []
    inventaire: list[EtatRenseigne] = []
    est_un_classeur = False

    if octets.startswith(b"%PDF"):
        normes, avertissements, pages = _analyser_pdf(octets)
        avertissements.append(
            "Lecture d'une impression : les contrôles de cohérence interne "
            "demandent le classeur. Sommer des nombres extraits d'un PDF ne "
            "prouverait rien, faute de savoir à quelle ligne du formulaire "
            "chacun appartient."
        )
    elif octets.startswith(b"PK"):
        # Un .xlsx est une archive ZIP : c'est la signature qui le dit.
        normes, controles, inventaire = analyser_classeur(octets)
        avertissements = []
        pages = 0
        est_un_classeur = True
    else:
        raise ValueError(
            "Format non reconnu : déposez le classeur de la déclaration (.xlsx) "
            "ou son impression (.pdf)."
        )

    if not normes:
        avertissements.append(
            "Aucune norme de l'EP01 n'a été trouvée. Ce document n'est pas une "
            "déclaration FODEP, ou son état de conformité est absent."
        )

    if (ecarts := [c for c in controles if c.statut == ECART]):
        avertissements.append(
            f"{len(ecarts)} total(aux) ne suit pas la somme de ses lignes, "
            "au-delà de ce que les arrondis expliquent. Un total qui contredit "
            "ses lignes fait rejeter le dépôt sans qu'aucune norme ne soit en "
            "cause."
        )

    if (vides := [e for e in inventaire if e.tout_a_zero]):
        avertissements.append(
            f"{len(vides)} état(s) sont renseignés entièrement à zéro. Le "
            "formulaire les déclare comme les autres : rien ne distingue un "
            "encours nul d'un encours qui n'a pas été mesuré."
        )

    manquantes = sorted(set(libelles_des_normes()) - {norme.code for norme in normes})
    if normes and manquantes:
        avertissements.append(
            "Normes absentes du document : " + ", ".join(manquantes) + "."
        )

    # Un plafond dont le niveau observe est nul se declare conforme. Rien ne
    # distingue, dans le document, un encours reellement nul d'un encours que le
    # declarant n'a pas su mesurer : c'est la meme case a zero. Le lecteur doit
    # le savoir avant de conclure.
    nulles = [
        norme.code
        for norme in normes
        if norme.minimum is False and norme.observe == 0
    ]
    if nulles:
        avertissements.append(
            "Niveau observé nul pour " + ", ".join(nulles) + " : ces normes se "
            "déclarent conformes, mais un zéro peut venir d'un encours non "
            "mesuré autant que d'un encours réellement nul."
        )

    non_mesurees = [n.code for n in normes if n.situation == NON_MESUREE]
    if non_mesurees:
        avertissements.append(
            "Niveau observé absent pour " + ", ".join(non_mesurees) + " : la case "
            "est calculée par le formulaire, et sa valeur n'existe qu'à "
            "l'ouverture du classeur. Une norme sans niveau observé n'est pas "
            "conforme : elle est non mesurée."
        )

    return AnalyseDeclaration(
        nom_fichier=nom_fichier,
        pages=pages,
        normes=normes,
        controles=controles,
        inventaire=inventaire,
        classeur=est_un_classeur,
        avertissements=avertissements,
    )


def analyser_declaration_en_cours(date_arrete=None) -> AnalyseDeclaration:
    """Analyse la declaration que l'outil produirait aujourd'hui.

    `analyser_declaration` lit un fichier recu ; celle-ci regarde ce que
    l'application s'apprete a transmettre. Meme lecture, meme vocabulaire :
    elle passe par `analyser_classeur`, de sorte qu'une norme se lise a
    l'identique qu'elle vienne d'un depot ou du portefeuille en base -- deux
    lectures separees finiraient par diverger sur le meme seuil.

    S'y ajoute ce qu'un fichier recu ne peut pas porter : les reserves que le
    remplissage a formulees. Une norme franchie et la reserve qui l'explique
    se lisent alors au meme endroit.
    """

    from app.rapports.fodep.service import renseigner_classeur_fodep

    produit = renseigner_classeur_fodep(date_arrete)
    try:
        tampon = BytesIO()
        produit.classeur.save(tampon)
        octets = tampon.getvalue()
        reserves = list(produit.anomalies)
    finally:
        produit.classeur.close()

    normes, controles, inventaire = analyser_classeur(octets)
    return AnalyseDeclaration(
        nom_fichier="Declaration en cours",
        pages=0,
        normes=normes,
        controles=controles,
        inventaire=inventaire,
        classeur=True,
        reserves=reserves,
    )
