"""Contenu du FODEP tel que le classeur le presente.

L'interface sait produire un PDF de la declaration. Ce PDF doit etre le classeur
imprime, pas une mise en page inventee : memes lignes, memes colonnes, memes
fusions, memes intitules en gras. Ce module extrait donc la *forme* autant que
les valeurs — largeurs de colonnes, cellules fusionnees, alignements — pour que
le lecteur retrouve l'etat qu'il connait.

L'extraction vit cote serveur parce que la bibliotheque Excel du poste de travail
echoue a decoder ce formulaire — « Null check operator used on a null value » des
le decodage — alors qu'openpyxl le relit sans peine, puisqu'il vient de l'ecrire.
Le classeur reste la seule source : le PDF n'est jamais un second calcul.
"""

from __future__ import annotations

from datetime import date

from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from app.rapports.fodep.reserves import Reserve
from app.rapports.fodep.service import (
    nom_fichier_fodep,
    renseigner_classeur_fodep,
)

# Garde-fou de largeur, pas une regle de mise en page.
#
# Il valait 14, au temps ou le PDF imprimait tout sur une A4 paysage sans savoir
# se reduire : au-dela, les colonnes devenaient illisibles. Mais cinq etats
# portent du texte plus loin — l'ADPE jusqu'a la colonne 25, l'EP07 jusqu'a la
# 54e — et la declaration imprimee en perdait des colonnes entieres. C'est le
# PDF qui choisit maintenant son format et sa reduction ; ce plafond ne sert
# plus qu'a ne pas suivre un formatage aberrant a l'infini.
COLONNES_MAX = 60

# Largeur de colonne par defaut d'Excel, en caracteres : openpyxl ne renseigne
# `width` que pour les colonnes explicitement dimensionnees.
LARGEUR_PAR_DEFAUT = 8.43

# Separateur de milliers : espace fine insecable, comme le veut la typographie
# francaise et comme l'interface formate ses montants. Insecable pour qu'un
# montant ne se coupe pas en fin de ligne, fine pour ne pas trouer le nombre.
ESPACE_FINE = " "


class CelluleFodep(BaseModel):
    """Une cellule du formulaire, avec ce qu'il faut pour la redessiner."""

    texte: str = ""

    #: Le formulaire met ses intitules et ses totaux en gras.
    gras: bool = False

    #: Les montants et les ratios sont alignes a droite dans le classeur.
    droite: bool = False

    #: Texte centre dans sa cellule, comme les en-tetes de colonne.
    centre: bool = False

    #: Nombre de colonnes couvertes, reprises des fusions du formulaire. Les
    #: titres d'etat et les intitules de section en couvrent plusieurs.
    colonnes: int = 1

    #: Couleur de fond en RVB (« FFF2CC »). Le formulaire teinte les cases a
    #: renseigner et ses en-tetes ; sans elles, le lecteur ne distingue plus ce
    #: que l'etablissement a declare de ce que le formulaire lui demandait. Le
    #: blanc n'est pas transmis : c'est deja la couleur du papier.
    fond: str | None = None

    #: Cotes bordes, parmi « l », « r », « t » et « b ». Le formulaire n'encadre
    #: que ses tableaux : border toutes les cellules donnait une grille uniforme
    #: qui ne ressemblait a aucune page du classeur. C'est aussi ce qui rend les
    #: fusions verticales continues, leur trait ne courant qu'en haut de la
    #: premiere cellule et en bas de la derniere.
    bordures: str = ""

    #: Taille de police du classeur : 10 pour le corps, jusqu'a 20 pour les
    #: titres. C'est elle qui donne au document sa hierarchie.
    taille: float | None = None


class LigneFodep(BaseModel):
    cellules: list[CelluleFodep] = Field(default_factory=list)

    #: Hauteur reglee dans le classeur, en points. Le formulaire aere ses
    #: lignes et donne a ses en-tetes deux ou trois fois la hauteur d'une
    #: ligne de donnees ; sans elle, toutes se serrent a la taille de leur
    #: texte et la page perd la forme du formulaire.
    hauteur: float | None = None

    #: Ligne d'en-tete, a reimprimer en haut de chaque page. Le classeur les
    #: designe lui-meme : c'est son reglage d'impression (« lignes a repeter en
    #: haut »). Un etat de cent lignes se lit autrement quand sa deuxieme page
    #: rappelle de quel etat et de quelles colonnes il s'agit.
    entete: bool = False


class EtatFodep(BaseModel):
    """Un etat du formulaire, reduit a ses lignes porteuses."""

    nom: str

    #: Largeurs des colonnes, dans l'unite d'Excel. Les conserver evite un PDF
    #: aux colonnes egales, ou l'intitule d'un poste serait aussi etroit que la
    #: colonne d'un code.
    largeurs: list[float] = Field(default_factory=list)

    #: Orientation reglee dans le classeur. Dix-neuf etats du FODEP sont en
    #: portrait : les imprimer tous en paysage etirait leurs colonnes sur une
    #: page trois fois trop large pour eux.
    paysage: bool = True

    lignes: list[LigneFodep] = Field(default_factory=list)


class ContenuFodep(BaseModel):
    """Le classeur rendu imprimable, avec ce qui doit voyager avec lui."""

    nom_fichier: str
    date_arrete: date

    #: Reserves de lecture de l'export, chacune avec sa nature. Elles ne
    #: figurent pas dans le classeur : l'interface les affiche, le PDF n'a pas
    #: a les inventer.
    anomalies: list[Reserve]

    etats: list[EtatFodep]


def _texte(valeur) -> str:
    """Cellule rendue lisible, ou vide quand elle n'a rien a dire."""

    if valeur is None:
        return ""
    if isinstance(valeur, bool):
        return "OUI" if valeur else "NON"
    if isinstance(valeur, str):
        texte = valeur.strip()
        # Une formule n'a pas de valeur hors d'Excel : le classeur est ecrit
        # sans etre evalue. Afficher « =SI(...) » ferait passer une mecanique
        # pour un resultat ; la case reste vide, comme dans le classeur non
        # encore ouvert.
        return "" if texte.startswith("=") else texte
    if isinstance(valeur, date):
        return valeur.strftime("%d/%m/%Y")
    if isinstance(valeur, float):
        if valeur == int(valeur):
            return f"{int(valeur):,}".replace(",", ESPACE_FINE)
        return f"{valeur:.4f}".replace(".", ",")
    if isinstance(valeur, int):
        return f"{valeur:,}".replace(",", ESPACE_FINE)
    return str(valeur).strip()


def _derniere_colonne(feuille) -> int:
    """Derniere colonne portant une valeur ou un trait, plafonnee a COLONNES_MAX.

    Le trait compte autant que le texte : la derniere colonne d'un tableau peut
    n'etre qu'une case a renseigner, vide et bordee. S'arreter au dernier texte
    lui retirait son bord droit.
    """

    derniere = 0
    for ligne in feuille.iter_rows(max_col=COLONNES_MAX):
        for cellule in ligne:
            if _texte(cellule.value) or _bordures(cellule):
                derniere = max(derniere, cellule.column)
    return derniere


def _fusions(
    feuille, derniere: int
) -> tuple[dict[tuple[int, int], int], set[tuple[int, int]]]:
    """Portees des fusions, et cellules qu'elles recouvrent.

    Le formulaire fusionne ses titres d'etat et ses intitules de section sur
    toute la largeur. Sans cette information, un titre de cent trente caracteres
    se retrouve comprime dans la colonne des codes DISPRU — et une colonne plus
    etroite qu'un seul de ses mots ne peut pas se mettre en page.
    """

    portees: dict[tuple[int, int], int] = {}
    recouvertes: set[tuple[int, int]] = set()
    for plage in feuille.merged_cells.ranges:
        if plage.min_col > derniere:
            continue
        largeur = min(plage.max_col, derniere) - plage.min_col + 1
        for rang in range(plage.min_row, plage.max_row + 1):
            portees[(rang, plage.min_col)] = largeur
            for colonne in range(plage.min_col + 1, plage.max_col + 1):
                recouvertes.add((rang, colonne))
    return portees, recouvertes


def _fond(cellule) -> str | None:
    """Couleur de remplissage, en RVB, ou rien si la cellule n'en porte pas."""

    remplissage = cellule.fill
    if remplissage is None or not remplissage.patternType:
        return None
    rvb = getattr(remplissage.fgColor, "rgb", None)
    if not isinstance(rvb, str) or len(rvb) not in (6, 8):
        return None
    couleur = rvb[-6:].upper()
    return None if couleur == "FFFFFF" else couleur


def _bordures(cellule) -> str:
    """Cotes bordes de la cellule, dans l'ordre gauche, droite, haut, bas."""

    bordure = cellule.border
    if bordure is None:
        return ""
    return "".join(
        cote[0]
        for cote in ("left", "right", "top", "bottom")
        if getattr(bordure, cote).style
    )


def _porte_quelque_chose(cellule: CelluleFodep) -> bool:
    """La cellule a-t-elle de quoi etre dessinee ?

    Une cellule sans texte compte quand meme si elle est bordee ou teintee :
    c'est une case a renseigner restee vide, et le tableau qui la contient
    perdrait sa forme si on la retirait.
    """

    return bool(cellule.texte or cellule.bordures or cellule.fond)


def _est_a_droite(cellule) -> bool:
    if cellule.alignment and cellule.alignment.horizontal == "right":
        return True
    # Un nombre sans alignement declare est cadre a droite par Excel.
    return isinstance(cellule.value, (int, float)) and not isinstance(
        cellule.value, bool
    )


def _lignes_de_titre(feuille) -> set[int]:
    """Lignes que le classeur reimprime en haut de chaque page.

    Le reglage est celui du formulaire lui-meme (« $1:$15 » pour l'EP30) : il
    dit ou finit l'en-tete et ou commence le tableau.
    """

    reglage = feuille.print_title_rows
    if not reglage:
        return set()
    try:
        debut, fin = reglage.replace("$", "").split(":")
        return set(range(int(debut), int(fin) + 1))
    except (ValueError, AttributeError):
        return set()


def _lignes_du_formulaire(feuille, derniere: int) -> list[LigneFodep]:
    portees, recouvertes = _fusions(feuille, derniere)
    titres = _lignes_de_titre(feuille)
    lignes: list[LigneFodep] = []

    for ligne in feuille.iter_rows(max_col=derniere):
        cellules: list[CelluleFodep] = []
        for cellule in ligne:
            position = (cellule.row, cellule.column)
            if position in recouvertes:
                continue
            cellules.append(
                CelluleFodep(
                    texte=_texte(cellule.value),
                    gras=bool(cellule.font and cellule.font.bold),
                    droite=_est_a_droite(cellule),
                    centre=bool(
                        cellule.alignment
                        and cellule.alignment.horizontal == "center"
                    ),
                    colonnes=portees.get(position, 1),
                    fond=_fond(cellule),
                    bordures=_bordures(cellule),
                    taille=float(cellule.font.sz)
                    if cellule.font and cellule.font.sz
                    else None,
                )
            )

        # Le formulaire compte beaucoup de lignes vides, reservees a la mise en
        # page : les reporter donnerait un PDF de blancs. Une ligne sans texte
        # mais bordee, elle, appartient a un tableau et se garde.
        if not any(_porte_quelque_chose(cellule) for cellule in cellules):
            continue
        # Les colonnes de queue qui ne portent rien du tout n'elargissent le
        # tableau pour rien.
        while cellules and not _porte_quelque_chose(cellules[-1]):
            cellules.pop()
        rang = ligne[0].row
        dimension = feuille.row_dimensions.get(rang)
        lignes.append(
            LigneFodep(
                cellules=cellules,
                hauteur=float(dimension.height)
                if dimension is not None and dimension.height
                else None,
                entete=rang in titres,
            )
        )

    return lignes


def _largeurs(feuille, derniere: int) -> list[float]:
    largeurs: list[float] = []
    for colonne in range(1, derniere + 1):
        dimension = feuille.column_dimensions.get(get_column_letter(colonne))
        largeur = dimension.width if dimension and dimension.width else None
        largeurs.append(float(largeur or LARGEUR_PAR_DEFAUT))
    return largeurs


def contenu_fodep(date_arrete: date | None = None) -> ContenuFodep:
    """Renseigne le FODEP puis en extrait le contenu, etat par etat.

    Le classeur est lu la ou il est produit, en memoire. L'ecrire dans un
    tampon pour le relire aussitot coutait huit secondes par apercu et ne
    rendait rien de plus : le classeur enregistre est le meme que celui qu'on
    tenait deja.
    """

    produit = renseigner_classeur_fodep(date_arrete)
    try:
        classeur = produit.classeur
        etats: list[EtatFodep] = []
        for nom in classeur.sheetnames:
            feuille = classeur[nom]
            derniere = _derniere_colonne(feuille)
            if derniere == 0:
                continue
            lignes = _lignes_du_formulaire(feuille, derniere)
            if lignes:
                etats.append(
                    EtatFodep(
                        nom=nom,
                        largeurs=_largeurs(feuille, derniere),
                        paysage=feuille.page_setup.orientation != "portrait",
                        lignes=lignes,
                    )
                )

        return ContenuFodep(
            nom_fichier=nom_fichier_fodep(produit.date_arrete),
            date_arrete=produit.date_arrete,
            anomalies=list(produit.anomalies),
            etats=etats,
        )
    finally:
        produit.classeur.close()
